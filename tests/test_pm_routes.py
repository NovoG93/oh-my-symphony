from __future__ import annotations

import base64
import asyncio
import os
from pathlib import Path
import shutil
import subprocess
import json
import hashlib
import difflib
import http.client
import threading
import time
import uuid
import io
from typing import Any

import pytest
import openpyxl
import yaml
from aiohttp.test_utils import TestClient, TestServer

from symphony.orchestrator import Orchestrator
from symphony.orchestrator.run_registry import RunRegistry
import symphony.pm.routes as routes_module
from symphony.server import build_app
from symphony.workflow import WorkflowState
from symphony.trackers.file import FileBoardTracker
from symphony.pm.documents import PmDocumentRepository
from symphony.pm.models import PmRequestContext
from symphony.pm.planning import PlanningProposal
from symphony.pm.routes import (
    PM_DEVELOPMENT_VERIFICATION_KEY,
    PM_DEMO_KEY,
    _apply_unified_diff,
    _assert_planning_reservations,
    _import_recovery_matches,
    _reserve_planning_ids,
)
from tests.test_pm_approval import _populate
from tests.test_pm_documents import _repository


@pytest.fixture(autouse=True)
def _isolate_host_context(tmp_path, monkeypatch):
    # Route fixtures are standalone projects, not entries in the user's registry.
    monkeypatch.setenv("SYMPHONY_PROJECTS_FILE", str(tmp_path / "projects.json"))
    for key in tuple(os.environ):
        if key in {
            "SYMPHONY_LANG", "SYMPHONY_API_TOKEN", "SYMPHONY_TRUSTED_ORIGINS",
            "SYMPHONY_SERVICE_INSTANCE_ID", "GIT_CONFIG_GLOBAL",
        } or key.startswith(("GIT_AUTHOR_", "GIT_COMMITTER_")):
            monkeypatch.delenv(key)
    # Assistance-on fixtures must never spawn the developer's host Codex.
    monkeypatch.setattr(routes_module.host_tools, "probe_codex", lambda: {
        "required": "restricted app-server controls", "found": "0.162.1", "available": True, "reason": None,
    })


async def _client(
    tmp_path: Path, *, authoring: bool = False, demo: bool = False,
    development_verification: bool = False,
) -> TestClient:
    (tmp_path / "kanban").mkdir()
    (tmp_path / "WORKFLOW.md").write_text(
        f"""---
tracker:
  kind: file
  board_root: ./kanban
  project_slug: p1
  active_states: [Todo, Doing]
  terminal_states: [Done]
agent:
  kind: claude
pm:
  authoring: {str(authoring).lower()}
  exchange: false
  assistance: false
  handoff: false
---
Prompt
""",
        encoding="utf-8",
    )
    state = WorkflowState(tmp_path / "WORKFLOW.md")
    cfg, error = state.reload()
    assert cfg is not None and error is None
    app = build_app(Orchestrator(state))
    if demo:
        app[PM_DEMO_KEY] = tmp_path.resolve()
    if development_verification:
        app[PM_DEVELOPMENT_VERIFICATION_KEY] = True
    client = TestClient(TestServer(app))
    await client.start_server()
    return client


async def _authoring_client(tmp_path: Path) -> tuple[TestClient, Path]:
    project = tmp_path / "project"
    project.mkdir()
    subprocess.run(["git", "-C", str(project), "init", "-q"], check=True)
    subprocess.run(["git", "-C", str(project), "config", "user.name", "test"], check=True)
    subprocess.run(["git", "-C", str(project), "config", "user.email", "test@example.com"], check=True)
    (project / "README.md").write_text("fixture\n", encoding="utf-8")
    subprocess.run(["git", "-C", str(project), "add", "."], check=True)
    subprocess.run(["git", "-C", str(project), "commit", "-qm", "fixture"], check=True)
    planning_root = tmp_path / "planning-root"
    planning_root.mkdir()
    planning = planning_root / "change"
    subprocess.run(["git", "-C", str(project), "worktree", "add", "-q", "-b", "planning", str(planning)], check=True)
    (project / "kanban").mkdir()
    (project / "WORKFLOW.md").write_text(
        f"""---
tracker:
  kind: file
  board_root: ./kanban
  project_slug: p1
  active_states: [Todo]
  terminal_states: [Done]
agent:
  kind: claude
pm:
  authoring: true
  exchange: true
  assistance: true
  handoff: true
  planning_workspace_root: {planning_root}
---
Prompt
""",
        encoding="utf-8",
    )
    state = WorkflowState(project / "WORKFLOW.md")
    cfg, error = state.reload()
    assert cfg is not None and error is None
    app = build_app(Orchestrator(state))
    app[PM_DEVELOPMENT_VERIFICATION_KEY] = True
    client = TestClient(TestServer(app))
    await client.start_server()
    return client, planning


async def _published_export_fixture(tmp_path: Path, monkeypatch, *, historical=False):
    from symphony.pm.models import ValidationReport

    monkeypatch.setattr(routes_module, "_openspec_readiness", lambda *_args: ValidationReport(True))
    client, planning = await _authoring_client(tmp_path)
    registry = RunRegistry(tmp_path / "project/.symphony/state.db")
    repository = PmDocumentRepository(project_id="p1", project_root=tmp_path / "project",
                                     registry=registry, planning_workspace_root=planning.parent)
    change_id = uuid.uuid4().hex
    repository.claim_workspace(change_id, planning, "owner")
    _populate(repository, change_id, planning)
    path = f"/api/v1/pm/changes/{change_id}"
    headers = {"X-Symphony-PM-Owner": "owner"}
    try:
        for index in range(2 if historical else 1):
            if index:
                tasks = planning / "openspec/changes/approved/tasks.md"
                tasks.write_text(tasks.read_text().replace("Implement approved scope", "Implement newer approved scope"))
            response = await client.get(path + "?project_id=p1", headers=headers)
            assert response.status == 200, await response.text()
            change = await response.json()
            envelope = {"project_id": "p1", "change_id": change_id, "expected_revision": change["revision"]}
            response = await client.post(path + "/approvals", headers=headers, json={
                **envelope, "confirmation_token": change["approval_confirmation_token"], "reviewer_label": "reviewer",
            })
            assert response.status == 201, await response.text()
            response = await client.post(path + "/handoffs/preview", headers=headers, json=envelope)
            assert response.status == 200, await response.text()
            preview = await response.json()
            response = await client.post(path + "/handoffs", headers=headers, json={
                **envelope, "client_operation_id": f"export-handoff-{index}",
                "preview_token": preview["preview_token"], "confirmation_token": preview["preview_token"],
            })
            assert response.status == 201, await response.text()
            result = await response.json()
            identifier = result["ticket_identifiers"][0]
        # The real native file is modified as a fixture's execution state, without dispatch.
        ticket = tmp_path / "project/kanban" / f"{identifier}.md"
        text = ticket.read_text()
        ticket.write_text(text.replace("state: Todo", "state: Done", 1))
        return client, registry, planning, path, headers, identifier
    except BaseException:
        registry.close()
        await client.close()
        raise


@pytest.mark.asyncio
async def test_xlsx_export_projects_matching_native_state_and_link_once_off_loop(tmp_path, monkeypatch):
    client, registry, planning, path, headers, identifier = await _published_export_fixture(tmp_path, monkeypatch, historical=True)
    scans = []
    original = FileBoardTracker.scan_all
    monkeypatch.setattr(FileBoardTracker, "scan_all", lambda self: (scans.append(threading.get_ident()), original(self))[1])
    before = {item: item.read_bytes() for item in (tmp_path / "project/kanban").glob("*.md")}
    canonical = {item: item.read_bytes() for item in (planning / "openspec/changes/approved").rglob("*") if item.is_file() and item.name != ".symphony-pm.lock"}
    try:
        response = await client.post(path + "/exports", headers=headers, json={"project_id": "p1", "format": "xlsx"})
        assert response.status == 201, await response.text()
        export = await response.json()
        data = await (await client.get(f"/api/v1/pm/exports/{export['export_id']}/file")).read()
        workbook = openpyxl.load_workbook(io.BytesIO(data))
        try:
            assert workbook["WBS"]["L2"].value == "Done"
            assert dict(workbook["Read Me"].values)["Supplied: Done"] == 1
            assert workbook["Traceability"].max_row == 4
            assert [workbook["Traceability"].cell(row, 4).value for row in range(2, 5)] == [f"/api/v1/issues/{identifier}"] * 3
            assert all(cell.hyperlink is None and cell.data_type != "f" for row in workbook["Traceability"] for cell in row)
        finally:
            workbook.close()
        assert scans == [scans[0]] and scans[0] != threading.get_ident()
        assert {item: item.read_bytes() for item in before} == before
        assert {item: item.read_bytes() for item in canonical} == canonical
        assert len(registry.list_pm_managed_tickets("p1")) == 2
        assert registry._connect().execute("SELECT count(*) FROM runs").fetchone()[0] == 0
    finally:
        registry.close()
        await client.close()


@pytest.mark.parametrize("mismatch", ["missing", "change_id", "task_id", "task_revision", "handoff_id", "snapshot_id", "unpublished", "stale", "wrong_project", "wrong_change", "batch_unpublished", "batch_snapshot", "approval_scope", "no_mapping"])
@pytest.mark.asyncio
async def test_xlsx_export_omits_unmatched_runtime_state_and_link(tmp_path, monkeypatch, mismatch):
    client, registry, planning, path, headers, identifier = await _published_export_fixture(tmp_path, monkeypatch)
    ticket = tmp_path / "project/kanban" / f"{identifier}.md"
    scans = []
    original = FileBoardTracker.scan_all
    try:
        if mismatch == "missing":
            ticket.unlink()
        elif mismatch in {"change_id", "task_id", "task_revision", "handoff_id", "snapshot_id"}:
            _, front, body = ticket.read_text().split("---", 2)
            values = yaml.safe_load(front)
            values["symphony_pm"][mismatch] = "wrong"
            ticket.write_text("---\n" + yaml.safe_dump(values, sort_keys=False) + "---" + body)
        elif mismatch == "no_mapping":
            registry._connect().execute("DELETE FROM pm_managed_tickets")
        elif mismatch in {"batch_unpublished", "batch_snapshot", "approval_scope"}:
            table, field, value = {
                "batch_unpublished": ("pm_handoff_batches", "state", "STAGING"),
                "batch_snapshot": ("pm_handoff_batches", "snapshot_id", "wrong"),
                "approval_scope": ("pm_approvals", "scope_hash", "wrong"),
            }[mismatch]
            registry._connect().execute(f"UPDATE {table} SET {field}=?", (value,))
        else:
            field, value = {"unpublished": ("state", "STAGING"), "stale": ("task_revision", "wrong"), "wrong_project": ("project_id", "p2"), "wrong_change": ("change_id", "wrong")}[mismatch]
            registry._connect().execute(f"UPDATE pm_managed_tickets SET {field}=?", (value,))
        monkeypatch.setattr(FileBoardTracker, "scan_all", lambda self: (scans.append(threading.get_ident()), original(self))[1])
        response = await client.post(path + "/exports", headers=headers, json={"project_id": "p1", "format": "xlsx"})
        assert response.status == 201, await response.text()
        export = await response.json()
        data = await (await client.get(f"/api/v1/pm/exports/{export['export_id']}/file")).read()
        workbook = openpyxl.load_workbook(io.BytesIO(data))
        try:
            assert workbook["WBS"]["L2"].value is None
            assert dict(workbook["Read Me"].values)["Not supplied"] == 1
            assert all(workbook["Traceability"].cell(row, 4).value is None for row in range(2, 5))
        finally:
            workbook.close()
        if mismatch in {"missing", "change_id", "task_id", "task_revision", "handoff_id", "snapshot_id"}:
            assert len(scans) == 1 and scans[0] != threading.get_ident()
        else:
            assert scans == []
    finally:
        registry.close()
        await client.close()


@pytest.mark.parametrize("format_name", ["markdown", "html", "zip", "csv"])
@pytest.mark.asyncio
async def test_non_xlsx_export_does_not_scan_native_runtime(tmp_path, monkeypatch, format_name):
    client, registry, _planning, path, headers, _identifier = await _published_export_fixture(tmp_path, monkeypatch)
    def unexpected_scan(_self):
        pytest.fail("non-XLSX export scanned native tickets")
    monkeypatch.setattr(FileBoardTracker, "scan_all", unexpected_scan)
    try:
        response = await client.post(path + "/exports", headers=headers, json={"project_id": "p1", "format": format_name})
        assert response.status == 201, await response.text()
    finally:
        registry.close()
        await client.close()


@pytest.mark.asyncio
async def test_capabilities_are_default_off_and_mutation_returns_stable_error(tmp_path: Path) -> None:
    client = await _client(tmp_path)
    try:
        response = await client.get("/api/v1/pm/capabilities")
        assert response.status == 200
        payload = await response.json()
        assert payload["flags"] == {
            "authoring": False, "exchange": False,
            "assistance": False, "handoff": False,
        }
        assert PM_DEVELOPMENT_VERIFICATION_KEY not in client.app
        assert PM_DEMO_KEY not in client.app
        assert payload["demo"] == {"active": False}
        assert payload["host_tools"] == {}
        change_id = "a" * 32
        response = await client.post(
            f"/api/v1/pm/changes/{change_id}/workspace/claim",
            json={"project_id": "p1", "workspace_path": str(tmp_path / "planning")},
        )
        assert response.status == 503
        assert (await response.json())["error"]["code"] == "capability_disabled"
    finally:
        await client.close()


@pytest.mark.parametrize(("demo", "development_verification"), [
    (False, False), (True, False), (False, True), (True, True),
])
@pytest.mark.asyncio
async def test_demo_constructor_flag_is_separate_from_release_gate_bypass(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, demo: bool, development_verification: bool,
) -> None:
    node = {
        "required": "20.19.0", "found": None, "available": False,
        "reason": "Bundled OpenSpec needs Node.js 20.19.0 or newer on PATH (found: none)",
    }
    monkeypatch.setattr(routes_module.host_tools, "probe_node", lambda: node)
    client = await _client(
        tmp_path, authoring=True, demo=demo,
        development_verification=development_verification,
    )
    try:
        app = client.app
        assert (PM_DEMO_KEY in app) is demo
        assert (PM_DEVELOPMENT_VERIFICATION_KEY in app) is development_verification
        response = await client.get("/api/v1/pm/capabilities")
        assert response.status == 200
        payload = await response.json()
        assert payload["demo"] == {"active": demo}
        assert payload["host_tools"] == {"node": node}
        assert payload["health"]["development_verification"] is development_verification
        assert payload["highest_passed_gate"] is None
        assert all(not gate["passed"] for gate in payload["gates"].values())
        assert payload["effective_flags"]["authoring"] is development_verification
        assert payload["capabilities"]["assistance"]["reasons"] == [
            "Off in the demo: it would contact an AI model." if demo
            else "PM assistance is disabled by host configuration"
        ]
        if not development_verification:
            # U2: even a PM-enabled build_app has no release-gate bypass.
            mutation = await client.post("/api/v1/pm/changes", json={
                "project_id": "p1", "workspace_path": str(tmp_path / "planning"),
                "locator": "locked", "title": "Locked",
            })
            assert mutation.status == 503
            assert (await mutation.json())["error"]["code"] == "capability_not_ready"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_host_tool_payload_does_not_replace_gate_reasons_or_effective_flags(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    available = {"required": "20.19.0", "found": "v22.22.3", "available": True, "reason": None}
    unavailable = {
        "required": "20.19.0", "found": None, "available": False,
        "reason": "Bundled OpenSpec needs Node.js 20.19.0 or newer on PATH (found: none)",
    }
    monkeypatch.setattr(routes_module.host_tools, "probe_node", lambda: available)
    client = await _client(tmp_path, authoring=True)
    try:
        before = await (await client.get("/api/v1/pm/capabilities")).json()
        monkeypatch.setattr(routes_module.host_tools, "probe_node", lambda: unavailable)
        after = await (await client.get("/api/v1/pm/capabilities")).json()
        for key in ("flags", "effective_flags", "capabilities", "gates", "demo"):
            assert after[key] == before[key]
        assert before["host_tools"] == {"node": available}
        assert after["host_tools"] == {"node": unavailable}
        # P6's UI order uses these existing gate reasons, then the separate Node reason.
        reasons = after["capabilities"]["authoring"]["reasons"]
        assert reasons[0].startswith("Requires passed release gates:")
        assert unavailable["reason"] not in reasons
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_assistance_on_capabilities_include_stubbed_codex_and_node(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    node = {"required": "20.19.0", "found": "v22.22.3", "available": True, "reason": None}
    monkeypatch.setattr(routes_module.host_tools, "probe_node", lambda: node)
    client, _ = await _authoring_client(tmp_path)
    try:
        response = await client.get("/api/v1/pm/capabilities")
        assert response.status == 200
        payload = await response.json()
        assert payload["demo"] == {"active": False}
        assert payload["host_tools"] == {
            "codex": {"required": "restricted app-server controls", "found": "0.162.1", "available": True, "reason": None},
            "node": node,
        }
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_pm_routes_reuse_global_media_and_bearer_guards(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    client = await _client(tmp_path)
    monkeypatch.setenv("SYMPHONY_API_TOKEN", "secret")
    try:
        assert (await client.get("/api/v1/pm/capabilities")).status == 401
        assert (
            await client.get(
                "/api/v1/pm/capabilities",
                headers={"Authorization": "Bearer secret"},
            )
        ).status == 200
        response = await client.post(
            f"/api/v1/pm/changes/{'a' * 32}/workspace/claim",
            data="{}",
            headers={"Authorization": "Bearer secret", "Content-Type": "text/plain"},
        )
        assert response.status == 415
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_registered_same_slug_routes_use_host_project_identity(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root_a = tmp_path / "a"
    root_b = tmp_path / "b"
    root_a.mkdir()
    root_b.mkdir()
    projects = tmp_path / "projects.json"
    projects.write_text(json.dumps({"version": 1, "projects": [
        {"id": "registered-a", "name": "A", "git_repo": str(root_a),
         "workflow": str(root_a / "WORKFLOW.md"), "host": "127.0.0.1", "port": 10001},
        {"id": "registered-b", "name": "B", "git_repo": str(root_b),
         "workflow": str(root_b / "WORKFLOW.md"), "host": "127.0.0.1", "port": 10002},
    ]}), encoding="utf-8")
    monkeypatch.setenv("SYMPHONY_PROJECTS_FILE", str(projects))
    client_a = await _client(root_a)
    client_b = await _client(root_b)
    try:
        assert (await (await client_a.get("/api/v1/pm/capabilities")).json())["project_id"] == "registered-a"
        assert (await (await client_b.get("/api/v1/pm/capabilities")).json())["project_id"] == "registered-b"
        stale = await client_a.get("/api/v1/pm/changes?project_id=registered-b")
        assert stale.status == 409
        assert (await stale.json())["error"]["code"] == "wrong_project"
    finally:
        await client_a.close()
        await client_b.close()


@pytest.mark.timing
@pytest.mark.asyncio
async def test_acceptance_post_refuses_cross_project_and_keeps_event_loop_live(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    started = threading.Event()
    release = threading.Event()

    def fake_preview(_self, **kwargs):
        return kwargs

    def blocking_accept(_self, **_kwargs):
        started.set()
        assert release.wait(5)
        return {"acceptance_id": "accepted"}

    monkeypatch.setattr(routes_module.PmVerificationService, "preview_acceptance", fake_preview)
    monkeypatch.setattr(routes_module.PmVerificationService, "accept", blocking_accept)
    client = await _client(tmp_path)
    try:
        wrong = await client.post(
            "/api/v1/pm/tickets/PM-1/acceptance",
            json={
                "project_id": "another-project", "task_revision": "r1",
                "tested_commit": "a" * 40, "confirmation_token": "token",
            },
        )
        assert wrong.status == 409
        assert (await wrong.json())["error"]["code"] == "wrong_project"
        accepting = asyncio.create_task(client.post(
            "/api/v1/pm/tickets/PM-1/acceptance",
            json={
                "project_id": "p1", "task_revision": "r1",
                "tested_commit": "a" * 40, "confirmation_token": "token",
            },
        ))
        for _ in range(100):
            if started.is_set():
                break
            await asyncio.sleep(0.01)
        assert started.is_set()
        heartbeat = await asyncio.wait_for(
            client.get("/api/v1/pm/capabilities"), timeout=1
        )
        assert heartbeat.status == 200
        release.set()
        accepted = await accepting
        assert accepted.status == 201
    finally:
        release.set()
        await client.close()


@pytest.mark.timing
@pytest.mark.asyncio
async def test_xlsx_export_and_import_preview_keep_state_route_responsive(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Loading a 1000-task change must not stall the event loop.

    Mirrors the release benchmark's api_exchange, whose heartbeat budget is
    0.25 s: /api/v1/state is polled while the export and the import preview
    run. With the change read on the loop the worst poll took 0.12-0.15 s;
    off the loop it takes 0.03-0.05 s, which is GIL handoff while workbook
    threads run. Wall-clock margins that narrow flap under host load, so the
    cause is asserted directly (no change read on the loop thread) and the
    latency only against the release budget.
    """
    # Keep the developer's host project registry out of the measured path.
    monkeypatch.setenv("SYMPHONY_PROJECTS_FILE", str(tmp_path / "projects.json"))
    loop_thread = threading.get_ident()
    read_threads: list[int] = []
    original_get_change = PmDocumentRepository.get_change

    def recording_get_change(self, *args, **kwargs):
        read_threads.append(threading.get_ident())
        return original_get_change(self, *args, **kwargs)

    monkeypatch.setattr(PmDocumentRepository, "get_change", recording_get_change)
    client, planning = await _authoring_client(tmp_path)
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "exchange-load", "title": "Exchange load",
        })
        assert created.status == 201, await created.text()
        change = await created.json()
        change_id = change["change_id"]
        seeded = await client.patch(
            f"/api/v1/pm/changes/{change_id}/wbs",
            json={"project_id": "p1", "change_id": change_id,
                  "expected_revision": change["revision"],
                  "payload": {"operations": [{"op": "create_task", "task": {
                      "kind": "Feature", "title": f"Task {index}", "order": index,
                      "description": f"Implement task {index}.",
                      "acceptance": [f"Task {index} is observable."],
                      "verification": ["focused"],
                      "integration_impact": "Existing PM workspace.",
                      "executable": True,
                  }} for index in range(1000)]}},
        )
        assert seeded.status == 200, await seeded.text()
        read_threads.clear()

        async def worst_state_latency(request) -> tuple[float, Any]:
            # Poll from a separate thread like the benchmark's external
            # client: an in-loop poller would sleep through a stall unmeasured.
            stop = threading.Event()
            samples: list[tuple[int, float]] = []
            port = client.port
            assert port is not None

            def poll_state() -> None:
                connection = http.client.HTTPConnection(
                    str(client.host), port, timeout=5,
                )
                try:
                    while not stop.is_set():
                        started = time.perf_counter()
                        connection.request("GET", "/api/v1/state")
                        response = connection.getresponse()
                        response.read()
                        samples.append((response.status, time.perf_counter() - started))
                        stop.wait(0.02)
                finally:
                    connection.close()

            poller = asyncio.create_task(asyncio.to_thread(poll_state))
            try:
                result = await request
            finally:
                stop.set()
                await poller
            assert samples and {status for status, _ in samples} == {200}
            return max(latency for _, latency in samples), result

        export_latency, exported = await worst_state_latency(client.post(
            f"/api/v1/pm/changes/{change_id}/exports",
            json={"project_id": "p1", "format": "xlsx"},
        ))
        assert exported.status == 201, await exported.text()
        export_id = (await exported.json())["export_id"]
        workbook = await (await client.get(f"/api/v1/pm/exports/{export_id}/file")).read()
        import_latency, previewed = await worst_state_latency(client.post(
            f"/api/v1/pm/changes/{change_id}/imports/preview",
            json={"project_id": "p1", "payload": {
                "filename": "exchange-load.xlsx",
                "base64": base64.b64encode(workbook).decode(),
            }},
        ))
        assert previewed.status == 200, await previewed.text()
        assert len(read_threads) >= 2, "export and import preview must read the change"
        assert loop_thread not in read_threads, "a 1000-task change was read on the event loop"
        assert export_latency < 0.25, f"export stalled /api/v1/state for {export_latency:.3f}s"
        assert import_latency < 0.25, f"import preview stalled /api/v1/state for {import_latency:.3f}s"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_pm_import_route_has_narrow_body_limit_override(tmp_path: Path) -> None:
    client = await _client(tmp_path)
    change_id = "a" * 32
    path = f"/api/v1/pm/changes/{change_id}/imports/preview"
    try:
        larger_than_global = base64.b64encode(b"x" * 900_000).decode()
        response = await client.post(
            path,
            json={
                "project_id": "p1", "change_id": change_id,
                "payload": {"filename": "large.xlsx", "base64": larger_than_global},
            },
        )
        assert response.status == 503
        assert (await response.json())["error"]["code"] == "capability_disabled"

        malformed = await client.post(
            path,
            json={
                "project_id": "p1", "change_id": change_id,
                "payload": {"filename": "malformed.xlsx", "base64": "not valid base64!"},
            },
        )
        assert malformed.status == 400
        assert (await malformed.json())["error"]["code"] == "invalid_request"

        encoded_over_limit = await client.post(
            path,
            json={
                "project_id": "p1", "change_id": change_id,
                "payload": {"filename": "oversized.xlsx", "base64": "A" * (8 * 1024 * 1024 + 4)},
            },
        )
        assert encoded_over_limit.status == 413
        assert (await encoded_over_limit.json())["error"]["code"] == "payload_too_large"

        unrelated = await client.post(
            "/api/v1/refresh",
            json={"padding": "x" * 1_100_000},
        )
        assert unrelated.status == 413
    finally:
        await client.close()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "filename",
    [
        None,
        "",
        " plan.xlsx",
        "plan.xlsx ",
        "plan.xls",
        "plan.xlsm",
        "plan.xlsx.exe",
        "../plan.xlsx",
        r"..\plan.xlsx",
    ],
)
async def test_pm_import_requires_safe_xlsx_filename(
    tmp_path: Path, filename: str | None,
) -> None:
    client = await _client(tmp_path)
    try:
        payload: dict[str, object] = {
            "base64": base64.b64encode(b"not-a-workbook").decode()
        }
        if filename is not None:
            payload["filename"] = filename
        response = await client.post(
            f"/api/v1/pm/changes/{'a' * 32}/imports/preview",
            json={"project_id": "p1", "payload": payload},
        )
        assert response.status == 415
        assert (await response.json())["error"]["code"] == "unsupported_media_type"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_pm_import_accepts_case_insensitive_xlsx_filename(
    tmp_path: Path,
) -> None:
    client, _ = await _authoring_client(tmp_path)
    change_id = "a" * 32
    try:
        response = await client.post(
            f"/api/v1/pm/changes/{change_id}/imports/preview",
            json={
                "project_id": "p1",
                "change_id": change_id,
                "payload": {
                    "filename": "Reviewed-Plan.XLSX",
                    "base64": base64.b64encode(b"not-a-workbook").decode(),
                },
            },
        )
        assert response.status == 400
        assert (await response.json())["error"]["code"] == "unsafe_workbook"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_pm_import_rejects_decoded_workbook_over_limit_with_413(
    tmp_path: Path,
) -> None:
    client, _ = await _authoring_client(tmp_path)
    change_id = "a" * 32
    try:
        response = await client.post(
            f"/api/v1/pm/changes/{change_id}/imports/preview",
            json={
                "project_id": "p1", "change_id": change_id,
                "payload": {
                    "filename": "oversized.xlsx",
                    "base64": base64.b64encode(
                        b"x" * (5 * 1024 * 1024 + 1)
                    ).decode()
                },
            },
        )
        assert response.status == 413
        payload = await response.json()
        assert payload["error"]["code"] == "payload_too_large"
        assert payload["error"]["message"] == "decoded workbook exceeds limit"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_pm_workspace_registers_operation_archive_and_acceptance_surfaces(tmp_path: Path) -> None:
    client = await _client(tmp_path)
    try:
        resources = {resource.canonical for resource in client.app.router.resources()}
        assert "/api/v1/pm/operations/{operation_id}" in resources
        assert "/api/v1/pm/operations/{operation_id}/cancel" in resources
        assert "/api/v1/pm/changes/{change_id}/archive/preview" in resources
        assert "/api/v1/pm/changes/{change_id}/archive/{archive_id}" in resources
        assert "/api/v1/pm/changes/{change_id}/specs" in resources
        assert "/api/v1/pm/tickets/{ticket_identifier}/verification-receipts" in resources
        assert "/api/v1/pm/tickets/{ticket_identifier}/acceptance/preview" in resources
        assert "/api/v1/pm/tickets/{ticket_identifier}/acceptance" in resources
        assert "/api/v1/pm/changes/{change_id}/planning-actions/{operation_id}/apply" in resources
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_cancel_operation_validates_project_before_registry_mutation(
    tmp_path: Path,
) -> None:
    client = await _client(tmp_path)
    registry = RunRegistry(tmp_path / ".symphony" / "state.db")
    operation_id = "a" * 32
    try:
        registry.create_pm_operation(
            operation_id=operation_id, project_id="another-project",
            change_id="b" * 32, kind="planning_proposal",
            expected_revision="sha256:revision",
        )
        response = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert response.status == 404
        row = registry._connect().execute(
            "SELECT state FROM pm_operations WHERE operation_id = ?", (operation_id,)
        ).fetchone()
        assert row is not None and row["state"] == "pending"
    finally:
        registry.close()
        await client.close()


@pytest.mark.asyncio
async def test_cancel_operation_keeps_pending_when_cleanup_is_uncertain(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    class UncertainClient:
        def __init__(self) -> None:
            self.started = asyncio.Event()
            self.release = asyncio.Event()

        async def propose(self, *_args, **_kwargs):
            self.started.set()
            await self.release.wait()
            raise routes_module.PmError(
                "planning_cleanup_failed", "cleanup uncertain", status=503
            )

        async def cancel_proposal(self) -> None:
            await self.stop()

        async def stop(self) -> None:
            self.release.set()
            raise routes_module.PmError(
                "planning_cleanup_failed", "cleanup uncertain", status=503
            )

    fake = UncertainClient()
    monkeypatch.setattr(routes_module, "PmPlanningClient", lambda **_kwargs: fake)
    client, planning = await _authoring_client(tmp_path)
    operation_id = "c" * 32
    registry = RunRegistry(tmp_path / "project" / ".symphony" / "state.db")
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "cancel-cleanup", "title": "Cancel cleanup",
        })
        assert created.status == 201, await created.text()
        change = await created.json()
        proposal_request = asyncio.create_task(client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
            json={
                "project_id": "p1", "change_id": change["change_id"],
                "expected_revision": change["revision"],
                "client_operation_id": operation_id,
                "payload": {"document_id": "proposal", "instruction": "Review"},
            },
        ))
        await fake.started.wait()
        cancelled = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert cancelled.status == 503
        assert (await cancelled.json())["error"]["code"] == "planning_cleanup_failed"
        row = registry._connect().execute(
            "SELECT state FROM pm_operations WHERE operation_id = ?", (operation_id,)
        ).fetchone()
        assert row is not None and row["state"] == "pending"
        response = await proposal_request
        assert response.status == 503
    finally:
        registry.close()
        await client.close()


@pytest.mark.asyncio
async def test_cancel_fence_survives_until_proposal_request_finishes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    class RacingClient:
        def __init__(self) -> None:
            self.started = asyncio.Event()
            self.result_ready = asyncio.Event()
            self.stop_entered = asyncio.Event()
            self.cleanup_allowed = asyncio.Event()
            self.stop_lock = asyncio.Lock()

        async def propose(self, *_args, **_kwargs):
            self.started.set()
            await self.result_ready.wait()
            try:
                return PlanningProposal(
                    "planning.md", "safe patch", False, "", "thread", "turn"
                )
            finally:
                await self.stop()

        async def cancel_proposal(self) -> None:
            await self.stop()

        async def stop(self) -> None:
            async with self.stop_lock:
                if not self.stop_entered.is_set():
                    self.stop_entered.set()
                    await self.cleanup_allowed.wait()

    fake = RacingClient()
    monkeypatch.setattr(routes_module, "PmPlanningClient", lambda **_kwargs: fake)
    client, planning = await _authoring_client(tmp_path)
    operation_id = "d" * 32
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "cancel-race", "title": "Cancel race",
        })
        assert created.status == 201, await created.text()
        change = await created.json()
        proposal_request = asyncio.create_task(client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
            json={
                "project_id": "p1", "change_id": change["change_id"],
                "expected_revision": change["revision"],
                "client_operation_id": operation_id,
                "payload": {"document_id": "proposal", "instruction": "Review"},
            },
        ))
        await fake.started.wait()
        cancel_request = asyncio.create_task(
            client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        )
        await fake.stop_entered.wait()
        fake.result_ready.set()
        await asyncio.sleep(0)
        fake.cleanup_allowed.set()
        cancelled = await cancel_request
        assert cancelled.status == 200, await cancelled.text()
        assert (await cancelled.json())["state"] == "cancelled"
        proposal = await proposal_request
        assert proposal.status == 409
        assert (await proposal.json())["error"]["code"] == "operation_cancelled"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_failed_operation_can_retry_retained_owner_cleanup(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    class RetryCleanupClient:
        cleanup_uncertain = True
        proposal_running = False

        def __init__(self) -> None:
            self.stop_calls = 0

        async def propose(self, *_args, **_kwargs):
            raise routes_module.PmError("planning_failed", "failed", status=503)

        async def stop(self) -> None:
            self.stop_calls += 1
            if self.stop_calls == 1:
                raise routes_module.PmError(
                    "planning_cleanup_failed", "cleanup uncertain", status=503
                )
            self.cleanup_uncertain = False

    fake = RetryCleanupClient()
    monkeypatch.setattr(routes_module, "PmPlanningClient", lambda **_kwargs: fake)
    client, planning = await _authoring_client(tmp_path)
    operation_id = "e" * 32
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "cleanup-retry", "title": "Cleanup retry",
        })
        change = await created.json()
        proposal = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
            json={
                "project_id": "p1", "change_id": change["change_id"],
                "expected_revision": change["revision"],
                "client_operation_id": operation_id,
                "payload": {"document_id": "proposal", "instruction": "Review"},
            },
        )
        assert proposal.status == 503
        retried = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert retried.status == 200, await retried.text()
        assert (await retried.json())["state"] == "failed"
        assert fake.stop_calls == 2
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_uncertain_owner_blocks_new_planning_operation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    class UncertainOwner:
        cleanup_uncertain = True
        proposal_running = False

        async def propose(self, *_args, **_kwargs):
            raise routes_module.PmError("planning_failed", "failed", status=503)

        async def stop(self) -> None:
            raise routes_module.PmError(
                "planning_cleanup_failed", "cleanup uncertain", status=503
            )

    instances: list[UncertainOwner] = []

    def factory(**_kwargs):
        instance = UncertainOwner()
        instances.append(instance)
        return instance

    monkeypatch.setattr(routes_module, "PmPlanningClient", factory)
    client, planning = await _authoring_client(tmp_path)
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "owner-admission", "title": "Owner admission",
        })
        change = await created.json()

        async def propose(operation_id: str):
            return await client.post(
                f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
                json={
                    "project_id": "p1", "change_id": change["change_id"],
                    "expected_revision": change["revision"],
                    "client_operation_id": operation_id,
                    "payload": {"document_id": "proposal", "instruction": "Review"},
                },
            )

        first = await propose("f" * 32)
        assert first.status == 503
        second = await propose("1" * 32)
        assert second.status == 503
        assert (await second.json())["error"]["code"] == "planning_cleanup_failed"
        assert len(instances) == 1
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_public_create_mints_change_id_and_unknown_claim_is_rejected(
    tmp_path: Path,
) -> None:
    client, planning = await _authoring_client(tmp_path)
    try:
        supplied = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "change_id": "a" * 32,
            "workspace_path": str(planning), "locator": "client-id", "title": "Client",
        })
        assert supplied.status == 409
        assert (await supplied.json())["error"]["code"] == "server_minted_id_required"
        unknown = await client.post(
            f"/api/v1/pm/changes/{'b' * 32}/workspace/claim",
            json={"project_id": "p1", "workspace_path": str(planning)},
        )
        assert unknown.status == 409
        assert (await unknown.json())["error"]["code"] == "server_minted_id_required"
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "server-id", "title": "Server",
        })
        assert created.status == 201, await created.text()
        payload = await created.json()
        assert len(payload["change_id"]) == 32
        assert payload["change_id"] not in {"a" * 32, "b" * 32}
        change_id = payload["change_id"]
        revision = payload["revision"]
        group = await client.patch(
            f"/api/v1/pm/changes/{change_id}/wbs",
            json={"project_id": "p1", "change_id": change_id,
                  "expected_revision": revision,
                  "payload": {"operations": [{"op": "create_task", "task": {
                      "kind": "Group", "title": "API group", "order": 10,
                  }}]}},
        )
        assert group.status == 200, await group.text()
        group_payload = await group.json()
        assert group_payload["impacted_ids"][0].startswith("G-")
        revision = group_payload["revision"]
        task = await client.patch(
            f"/api/v1/pm/changes/{change_id}/wbs",
            json={"project_id": "p1", "change_id": change_id,
                  "expected_revision": revision,
                  "payload": {"operations": [{"op": "create_task", "task": {
                      "kind": "Feature", "title": "API task", "order": 20,
                  }}]}},
        )
        assert task.status == 200, await task.text()
        task_payload = await task.json()
        task_id = task_payload["impacted_ids"][0]
        assert task_id.startswith("T-")
        supplied_task = await client.patch(
            f"/api/v1/pm/changes/{change_id}/wbs",
            json={"project_id": "p1", "change_id": change_id,
                  "expected_revision": task_payload["revision"],
                  "payload": {"operations": [{"op": "add_task",
                      "task_id": "T-CLIENT01",
                      "task": {"task_id": "T-CLIENT01", "title": "Client"}}]}},
        )
        assert supplied_task.status == 409
        assert (await supplied_task.json())["error"]["code"] == "server_minted_id_required"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_archived_original_change_id_is_readonly_and_deep_linkable(
    tmp_path: Path,
) -> None:
    client, planning = await _authoring_client(tmp_path)
    registry = RunRegistry(tmp_path / "project" / ".symphony" / "state.db")
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "history-link", "title": "History link",
        })
        assert created.status == 201, await created.text()
        change = await created.json()
        change_id = change["change_id"]
        active = planning / "openspec" / "changes" / "history-link"
        archived = planning / "openspec" / "changes" / "archive" / "2026-10-05-history-link"
        archived.parent.mkdir(parents=True, exist_ok=True)
        active.rename(archived)
        claim = registry.get_pm_workspace("p1", change_id)
        assert claim is not None
        assert registry.set_pm_workspace_state(
            "p1", change_id, fence=claim.fence, state="RELEASED"
        )

        change_response = await client.get(
            f"/api/v1/pm/changes/{change_id}?project_id=p1"
        )
        assert change_response.status == 200, await change_response.text()
        history = await change_response.json()
        assert history["change_id"] == change_id
        assert history["state"] == "Archived"
        assert history["workspace_state"] == "RELEASED"
        assert history["approval_confirmation_token"] is None

        for suffix in (
            "documents/tasks", "baselines/requirements", "wbs",
            "managed-tickets",
        ):
            response = await client.get(
                f"/api/v1/pm/changes/{change_id}/{suffix}?project_id=p1"
            )
            assert response.status == 200, (suffix, await response.text())
        document = await (
            await client.get(
                f"/api/v1/pm/changes/{change_id}/documents/tasks?project_id=p1"
            )
        ).json()
        assert document["editable"] is False
        assert document["unsupported_reason"] == "archived planning history"

        listed = await (await client.get("/api/v1/pm/changes?project_id=p1")).json()
        assert all(item["change_id"] != change_id for item in listed["changes"])
        write = await client.patch(
            f"/api/v1/pm/changes/{change_id}/documents/proposal",
            json={
                "project_id": "p1", "change_id": change_id,
                "expected_revision": history["revision"],
                "payload": {"operations": [{
                    "op": "replace_source", "source": "# changed\n",
                }]},
            },
        )
        assert write.status == 409
        assert (await write.json())["error"]["code"] == "workspace_repair_required"
        handoff = await client.post(
            f"/api/v1/pm/changes/{change_id}/handoffs/preview",
            json={
                "project_id": "p1", "change_id": change_id,
                "expected_revision": history["revision"],
            },
        )
        assert handoff.status == 409
        assert (await handoff.json())["error"]["code"] == "workspace_repair_required"

        outside = tmp_path / "outside-proposal.md"
        outside.write_text("outside archive sentinel\n", encoding="utf-8")
        archived_proposal = archived / "proposal.md"
        archived_proposal.unlink()
        archived_proposal.symlink_to(outside)
        unsafe = await client.get(
            f"/api/v1/pm/changes/{change_id}?project_id=p1"
        )
        assert unsafe.status == 400
        assert (await unsafe.json())["error"]["code"] == "unsafe_path"

        duplicate = archived.parent / "2026-10-06-history-link-copy"
        shutil.copytree(archived, duplicate, symlinks=True)
        ambiguous = await client.get(
            f"/api/v1/pm/changes/{change_id}?project_id=p1"
        )
        assert ambiguous.status == 409
        assert (
            await ambiguous.json()
        )["error"]["code"] == "ambiguous_change_identity"
    finally:
        registry.close()
        await client.close()


def test_server_side_planning_diff_rejects_stale_context() -> None:
    patch = "@@ -1,2 +1,2 @@\n one\n-two\n+three\n"
    assert _apply_unified_diff("one\ntwo\n", patch) == "one\nthree\n"
    with pytest.raises(Exception, match="no longer matches"):
        _apply_unified_diff("changed\ntwo\n", patch)


def test_server_side_planning_diff_supports_bounded_line_edge_cases() -> None:
    assert _apply_unified_diff("", "@@ -0,0 +1 @@\n+one\n") == "one\n"
    assert _apply_unified_diff("one\n", "@@ -1 +0,0 @@\n-one\n") == ""
    assert _apply_unified_diff(
        "one", "@@ -1 +1 @@\n-one\n\\ No newline at end of file\n+two\n\\ No newline at end of file"
    ) == "two"
    assert _apply_unified_diff(
        "one", "@@ -1 +1 @@\n-one\n\\ No newline at end of file\n+two\n"
    ) == "two\n"
    with pytest.raises(Exception, match="unsupported unified diff line"):
        _apply_unified_diff("one\n", "@@ -1 +1 @@\n one\ntrailing junk")
    with pytest.raises(Exception, match="invalid zero line position"):
        _apply_unified_diff("one\n", "@@ -1 +0 @@\n-one\n+changed\n")
    with pytest.raises(Exception, match="new position is inconsistent"):
        _apply_unified_diff(
            "one\ntwo\nthree\nfour\n",
            "@@ -1 +1 @@\n-one\n+first\n"
            "@@ -3 +2 @@\n-three\n+third\n",
        )
    assert _apply_unified_diff(
        "one\ntwo\nthree\nfour\n",
        "@@ -1 +1 @@\n-one\n+first\n"
        "@@ -3 +3 @@\n-three\n+third\n",
    ) == "first\ntwo\nthird\nfour\n"
    assert _apply_unified_diff(
        "one\r\ntwo\r\n",
        "@@ -1,2 +1,2 @@\n one\n-two\n+three\n",
    ) == "one\r\nthree\r\n"


def test_import_crash_reconciliation_requires_exact_full_manifest(tmp_path: Path) -> None:
    repository, change_id, workspace = _repository(tmp_path)
    change = _populate(repository, change_id, workspace, "import-recovery")
    operations = [{"op": "set_field", "task_id": "T-ABCDEFGH", "field": "title", "value": "Imported"}]
    ctx = PmRequestContext("p1", change_id, change.revision, "import-operation")
    prepared = repository.preview_wbs_patch(ctx, operations, owner="owner")
    assert not _import_recovery_matches(repository, ctx, prepared, owner="owner")
    applied = repository.patch_wbs(ctx, operations, owner="owner")
    current = PmRequestContext("p1", change_id)
    assert _import_recovery_matches(repository, current, prepared, owner="owner")

    repository.patch_wbs(
        PmRequestContext("p1", change_id, applied.revision),
        [{"op": "set_field", "task_id": "T-ABCDEFGH", "field": "notes", "value": "unrelated drift"}],
        owner="owner",
    )
    assert not _import_recovery_matches(repository, current, prepared, owner="owner")


@pytest.mark.asyncio
async def test_manual_report_download_requires_stable_project_identity(tmp_path: Path) -> None:
    client, _planning = await _authoring_client(tmp_path)
    project = tmp_path / "project"
    registry = RunRegistry(project / ".symphony" / "state.db")
    try:
        report = project / "manual-report.json"
        report.write_text('{"result":"passed"}\n', encoding="utf-8")
        subprocess.run(["git", "-C", str(project), "add", "manual-report.json"], check=True)
        subprocess.run(["git", "-C", str(project), "commit", "-qm", "manual report"], check=True)
        commit = subprocess.check_output(["git", "-C", str(project), "rev-parse", "HEAD"], text=True).strip()
        payload = report.read_bytes()
        ref = {"ref_kind": "git_blob", "ref": "manual-report.json", "media_kind": "json", "sha256": "sha256:" + hashlib.sha256(payload).hexdigest(), "size_bytes": len(payload)}
        registry.record_pm_manual_verification({
            "receipt_id": "manual-receipt", "operation_id": "manual-op",
            "project_id": "p1", "change_id": "c1", "ticket_identifier": "PM-1",
            "task_id": "T-1", "task_revision": "r1", "procedure_id": "manual",
            "procedure_kind": "manual", "definition_digest": "sha256:def",
            "tested_commit": commit, "environment_summary": "browser", "return_code": None,
            "timed_out": 0, "cancelled": 0, "stdout_hash": "sha256:out", "stderr_hash": "sha256:err",
            "started_at": "2026-10-05T00:00:00Z", "completed_at": "2026-10-05T00:01:00Z",
            "submitted_at": "2026-10-05T00:01:00Z", "submitted_by": "reviewer",
            "report_ref_json": json.dumps(ref), "artifact_refs_json": "[]",
            "result": "passed", "input_hash": "sha256:input",
        })
        assert (await client.get("/api/v1/pm/evidence/manual-receipt/report")).status == 409
        downloaded = await client.get("/api/v1/pm/evidence/manual-receipt/report?project_id=p1")
        assert downloaded.status == 200
        assert await downloaded.read() == payload
    finally:
        registry.close()
        await client.close()


def test_planning_placeholder_ids_are_host_reserved_and_operation_stable(tmp_path: Path) -> None:
    registry = RunRegistry(tmp_path / "state.db")
    try:
        proposals = [{
            "document_id": "tasks",
            "proposal_patch": "@@SYMPHONY_NEW_G_GROUP@@ @@SYMPHONY_NEW_T_TASK@@",
        }]
        first, rewritten = _reserve_planning_ids(registry, "planning-op", proposals)
        second, replay = _reserve_planning_ids(registry, "planning-op", proposals)
        assert first == second
        assert rewritten == replay
        assert set(first.values()) == set(rewritten[0]["proposal_patch"].split())
        assert all(value.startswith(("G-", "T-")) for value in first.values())
        _assert_planning_reservations(registry, "planning-op", first)
        registry._connect().execute(
            "UPDATE pm_planning_reserved_ids SET state='RETIRED' WHERE operation_id=?",
            ("planning-op",),
        )
        with pytest.raises(routes_module.PmError) as exc_info:
            _assert_planning_reservations(registry, "planning-op", first)
        assert exc_info.value.code == "operation_not_ready"
    finally:
        registry.close()


def test_planning_placeholder_ids_reject_wrong_document(tmp_path: Path) -> None:
    registry = RunRegistry(tmp_path / "state.db")
    try:
        with pytest.raises(routes_module.PmError) as exc_info:
            _reserve_planning_ids(
                registry,
                "planning-op",
                [{"document_id": "proposal", "proposal_patch": "@@SYMPHONY_NEW_T_TASK@@"}],
            )
        assert exc_info.value.code == "invalid_proposal"
    finally:
        registry.close()


@pytest.mark.asyncio
async def test_propose_wbs_applies_useful_host_ids_before_change_is_ready(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    class WbsPlanningClient:
        async def propose(self, _cwd: Path, snapshot: str, _instruction: str):
            source = json.loads(snapshot)["documents"]["tasks"]
            addition = (
                "\n\n<!-- symphony-group:@@SYMPHONY_NEW_G_DELIVERY@@ -->\n"
                "## Delivery\n```yaml symphony-group\n"
                "id: @@SYMPHONY_NEW_G_DELIVERY@@\nparent_id: null\norder: 10\n"
                "```\n<!-- /symphony-group -->\n\n"
                "<!-- symphony-task:@@SYMPHONY_NEW_T_IMPLEMENT@@ -->\n"
                "- [ ] Implement the approved behavior\n```yaml symphony-task\n"
                "id: @@SYMPHONY_NEW_T_IMPLEMENT@@\n"
                "parent_id: @@SYMPHONY_NEW_G_DELIVERY@@\norder: 10\n"
                "type: Feature\ndepends_on: []\nrequirements: []\n"
                "storyboards: []\ndesign_refs: []\n```\n\n"
                "### Scope\nImplement the bounded behavior.\n\n"
                "### Acceptance criteria\n- The focused behavior is verified.\n"
                "<!-- /symphony-task -->\n"
            )
            updated = source.rstrip() + addition
            patch = "".join(difflib.unified_diff(
                source.splitlines(keepends=True), updated.splitlines(keepends=True),
                fromfile="a/tasks.md", tofile="b/tasks.md",
            ))
            return PlanningProposal(
                "tasks", patch, False, "", "thread", "turn",
                explanation="Proposed the first implementation task.",
                proposals=({"document_id": "tasks", "proposal_patch": patch},),
            )

        async def stop(self) -> None:
            return None

    monkeypatch.setattr(
        routes_module, "PmPlanningClient", lambda **_kwargs: WbsPlanningClient()
    )
    client, planning = await _authoring_client(tmp_path)
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "early-wbs", "title": "Early WBS",
        })
        assert created.status == 201, await created.text()
        change = await created.json()
        operation_id = "2" * 32
        proposed_response = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
            json={
                "project_id": "p1", "expected_revision": change["revision"],
                "client_operation_id": operation_id, "action": "propose_wbs",
                "instruction": "Create the first bounded task.",
                "impacted_document_ids": ["tasks"],
            },
        )
        assert proposed_response.status == 201, await proposed_response.text()
        proposed = await proposed_response.json()
        reserved = proposed["result"]["reserved_ids"]
        assert {value[:2] for value in reserved.values()} == {"G-", "T-"}
        applied_response = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions/{operation_id}/apply",
            json={
                "project_id": "p1", "expected_revision": change["revision"],
                "client_operation_id": "3" * 32,
                "proposal_digest": proposed["result"]["proposal_digest"],
            },
        )
        assert applied_response.status == 200, await applied_response.text()
        applied = await applied_response.json()
        assert applied["revision"] != change["revision"]
        assert applied["findings"]
        wbs_response = await client.get(
            f"/api/v1/pm/changes/{change['change_id']}/wbs?project_id=p1"
        )
        wbs = await wbs_response.json()
        assert {item["task_id"] for item in wbs["tasks"]} == set(reserved.values())
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_truncated_planning_diff_is_rejected_before_document_transaction(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    class TruncatedPlanningClient:
        async def propose(self, _cwd: Path, snapshot: str, _instruction: str):
            source = json.loads(snapshot)["documents"]["proposal"]
            first_three = source.splitlines()[:3]
            assert len(first_three) == 3
            patch = (
                "--- a/proposal.md\n+++ b/proposal.md\n"
                "@@ -1,3 +1,3 @@\n"
                f" {first_three[0]}\n-{first_three[1]}\n+changed\n"
            )
            return PlanningProposal(
                "proposal", patch, False, "", "thread", "turn",
                proposals=({"document_id": "proposal", "proposal_patch": patch},),
            )

        async def stop(self) -> None:
            return None

    monkeypatch.setattr(
        routes_module, "PmPlanningClient", lambda **_kwargs: TruncatedPlanningClient()
    )
    client, planning = await _authoring_client(tmp_path)
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "truncated-diff", "title": "Truncated diff",
        })
        change = await created.json()
        before = await (
            await client.get(
                f"/api/v1/pm/changes/{change['change_id']}/documents/proposal?project_id=p1"
            )
        ).json()
        response = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
            json={
                "project_id": "p1", "expected_revision": change["revision"],
                "client_operation_id": "f" * 32, "action": "draft_artifact",
                "document_id": "proposal", "instruction": "Revise the proposal.",
            },
        )
        assert response.status == 409, await response.text()
        assert (await response.json())["error"]["code"] == "invalid_proposal"
        after_change = await (
            await client.get(
                f"/api/v1/pm/changes/{change['change_id']}?project_id=p1"
            )
        ).json()
        after = await (
            await client.get(
                f"/api/v1/pm/changes/{change['change_id']}/documents/proposal?project_id=p1"
            )
        ).json()
        assert after_change["revision"] == change["revision"]
        assert after["source"] == before["source"]
        change_root = planning / "openspec" / "changes" / "truncated-diff"
        assert not (change_root / ".symphony-multi-document-journal.json").exists()
    finally:
        await client.close()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "malformed_kind",
    (
        "new_coordinate", "no_newline_marker", "context_append",
        "zero_context_append", "duplicate_header",
    ),
)
async def test_malformed_planning_coordinates_and_eof_markers_never_write(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    malformed_kind: str,
) -> None:
    class MalformedPlanningClient:
        async def propose(self, _cwd: Path, snapshot: str, _instruction: str):
            source = json.loads(snapshot)["documents"]["proposal"]
            lines = source.splitlines()
            assert len(lines) >= 2
            if malformed_kind == "new_coordinate":
                malformed_patch = (
                    f"@@ -1,1 +999,1 @@\n-{lines[0]}\n+changed\n"
                )
            elif malformed_kind == "no_newline_marker":
                malformed_patch = (
                    f"@@ -1,2 +1,2 @@\n {lines[0]}\n"
                    f"\\ No newline at end of file\n {lines[1]}\n"
                )
            elif malformed_kind == "context_append":
                line_number = len(lines)
                malformed_patch = (
                    f"@@ -{line_number},1 +{line_number},2 @@\n"
                    f" {lines[-1]}\n+joined\n"
                )
            elif malformed_kind == "zero_context_append":
                line_number = len(lines)
                malformed_patch = (
                    f"@@ -{line_number},0 +{line_number + 1},1 @@\n+joined\n"
                )
            else:
                malformed_patch = (
                    "+++ b/proposal.md\n+++ b/duplicate.md\n"
                    f"@@ -1 +1 @@\n-{lines[0]}\n+changed\n"
                )
            return PlanningProposal(
                "proposal", malformed_patch, False, "", "thread", "turn",
                proposals=({
                    "document_id": "proposal",
                    "proposal_patch": malformed_patch,
                },),
            )

        async def stop(self) -> None:
            return None

    monkeypatch.setattr(
        routes_module, "PmPlanningClient", lambda **_kwargs: MalformedPlanningClient()
    )
    client, planning = await _authoring_client(tmp_path)
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "malformed-diff", "title": "Malformed diff",
        })
        change = await created.json()
        proposal_path = (
            planning / "openspec" / "changes" / "malformed-diff" / "proposal.md"
        )
        proposal_path.write_text(
            proposal_path.read_text(encoding="utf-8").rstrip("\n"),
            encoding="utf-8",
        )
        change = await (
            await client.get(
                f"/api/v1/pm/changes/{change['change_id']}?project_id=p1"
            )
        ).json()
        document_url = (
            f"/api/v1/pm/changes/{change['change_id']}"
            "/documents/proposal?project_id=p1"
        )
        before = await (await client.get(document_url)).json()
        response = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
            json={
                "project_id": "p1", "expected_revision": change["revision"],
                "client_operation_id": uuid.uuid4().hex,
                "action": "draft_artifact", "document_id": "proposal",
                "instruction": "Revise the proposal.",
            },
        )
        assert response.status == 409, await response.text()
        assert (await response.json())["error"]["code"] == "invalid_proposal"
        after_change = await (
            await client.get(
                f"/api/v1/pm/changes/{change['change_id']}?project_id=p1"
            )
        ).json()
        after = await (await client.get(document_url)).json()
        assert after_change["revision"] == change["revision"]
        assert after["source"] == before["source"]
        change_root = planning / "openspec" / "changes" / "malformed-diff"
        assert not (change_root / ".symphony-multi-document-journal.json").exists()
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_propose_wbs_retries_unmanaged_checkbox_with_canonical_grammar(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    class CorrectingClient:
        calls = 0

        async def propose(self, _cwd: Path, snapshot: str, instruction: str):
            self.calls += 1
            source = json.loads(snapshot)["documents"]["tasks"]
            if self.calls == 1:
                assert "previous proposal" not in instruction
                updated = source.rstrip() + "\n\n- [ ] T-PLACEHOLDER **Feature: unsafe**\n"
            else:
                assert "previous proposal was structurally invalid" in instruction
                updated = source.rstrip() + (
                    "\n\n<!-- symphony-group:@@SYMPHONY_NEW_G_RETRY@@ -->\n"
                    "## Retry\n```yaml symphony-group\n"
                    "id: @@SYMPHONY_NEW_G_RETRY@@\nparent_id: null\norder: 10\n"
                    "```\n<!-- /symphony-group -->\n"
                )
            patch = "".join(difflib.unified_diff(
                source.splitlines(keepends=True), updated.splitlines(keepends=True),
                fromfile="a/tasks.md", tofile="b/tasks.md",
            ))
            return PlanningProposal(
                "tasks", patch, False, "", "thread", f"turn-{self.calls}",
                explanation="Corrected WBS grammar.",
                proposals=({"document_id": "tasks", "proposal_patch": patch},),
            )

        async def stop(self) -> None:
            return None

    fake = CorrectingClient()
    monkeypatch.setattr(routes_module, "PmPlanningClient", lambda **_kwargs: fake)
    client, planning = await _authoring_client(tmp_path)
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "retry-wbs", "title": "Retry WBS",
        })
        change = await created.json()
        response = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
            json={
                "project_id": "p1", "expected_revision": change["revision"],
                "client_operation_id": "6" * 32, "action": "propose_wbs",
                "instruction": "Create a managed group.",
                "impacted_document_ids": ["tasks"],
            },
        )
        assert response.status == 201, await response.text()
        result = (await response.json())["result"]
        assert fake.calls == 2
        assert list(result["reserved_ids"].values())[0].startswith("G-")
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_all_seven_planning_actions_validate_server_owned_scope(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    class ExplanationClient:
        async def propose(self, _cwd: Path, _snapshot: str, instruction: str):
            return PlanningProposal(
                "", "", False, "", "thread", "turn",
                explanation=f"Explained {instruction}", proposals=(),
            )

        async def stop(self) -> None:
            return None

    monkeypatch.setattr(
        routes_module, "PmPlanningClient", lambda **_kwargs: ExplanationClient()
    )
    client, planning = await _authoring_client(tmp_path)
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "seven-actions", "title": "Seven actions",
        })
        change = await created.json()
        revision = change["revision"]
        task_ids: list[str] = []
        for title in ("First", "Second"):
            response = await client.patch(
                f"/api/v1/pm/changes/{change['change_id']}/wbs",
                json={
                    "project_id": "p1", "expected_revision": revision,
                    "payload": {"operations": [{"op": "create_task", "task": {
                        "kind": "Feature", "title": title,
                    }}]},
                },
            )
            assert response.status == 200, await response.text()
            value = await response.json()
            task_ids.append(value["impacted_ids"][0])
            revision = value["revision"]
        dependency = await client.patch(
            f"/api/v1/pm/changes/{change['change_id']}/wbs",
            json={
                "project_id": "p1", "expected_revision": revision,
                "payload": {"operations": [{
                    "op": "set_field", "task_id": task_ids[1],
                    "field": "depends_on",
                    "value": [{"task_id": task_ids[0], "kind": "requires_code"}],
                }]},
            },
        )
        assert dependency.status == 200, await dependency.text()
        revision = (await dependency.json())["revision"]
        current = await (
            await client.get(
                f"/api/v1/pm/changes/{change['change_id']}?project_id=p1"
            )
        ).json()
        section = next(
            item for item in current["planning_sections"]
            if item["document_id"] == "proposal"
        )
        requests = (
            {"action": "draft_artifact", "document_id": "proposal"},
            {"action": "refine_section", "document_id": "proposal", "section_id": section["section_id"]},
            {"action": "find_gaps", "impacted_document_ids": ["proposal"]},
            {"action": "review_consistency", "impacted_document_ids": ["proposal", "ux"]},
            {"action": "explain_dependency", "dependency": {
                "task_id": task_ids[1], "predecessor_id": task_ids[0],
            }},
            {"action": "propose_wbs"},
            {"action": "update_impacted_artifacts", "document_id": "proposal", "impacted_document_ids": ["proposal", "ux"]},
        )
        for index, action_payload in enumerate(requests, start=1):
            response = await client.post(
                f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
                json={
                    "project_id": "p1", "expected_revision": revision,
                    "client_operation_id": f"{index:032x}",
                    "instruction": "Review the selected scope.",
                    **action_payload,
                },
            )
            assert response.status == 201, (
                action_payload["action"], await response.text()
            )
            result = (await response.json())["result"]
            assert result["action"] == action_payload["action"]
            assert result["proposals"] == []
            assert result["write_or_dispatch_performed"] is False
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_coverage_waiver_is_revision_bound_idempotent_and_non_authoritative(
    tmp_path: Path,
) -> None:
    client, planning = await _authoring_client(tmp_path)
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(planning),
            "locator": "coverage-waiver", "title": "Coverage waiver",
        })
        change = await created.json()
        scenario_id = next(
            item["semantic_id"] for item in change["semantic_index"]
            if item["kind"] == "scenario"
        )
        preview_response = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/coverage/waivers/preview",
            json={
                "project_id": "p1", "expected_revision": change["revision"],
                "semantic_id": scenario_id, "reason": "Hardware-only path",
                "reviewer_label": "reviewer",
            },
        )
        assert preview_response.status == 200, await preview_response.text()
        preview = await preview_response.json()
        create_payload = {
            "project_id": "p1", "expected_revision": change["revision"],
            "client_operation_id": "4" * 32, "semantic_id": scenario_id,
            "reason": "Hardware-only path", "reviewer_label": "reviewer",
            "confirmation_token": preview["confirmation_token"],
        }
        first = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/coverage/waivers",
            json=create_payload,
        )
        assert first.status == 201, await first.text()
        waiver = await first.json()
        replay = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/coverage/waivers",
            json=create_payload,
        )
        assert replay.status == 200
        coverage = await (
            await client.get(
                f"/api/v1/pm/changes/{change['change_id']}/coverage"
                f"?project_id=p1&revision={change['revision']}"
            )
        ).json()
        entity = next(
            item for item in coverage["entities"]
            if item["semantic_id"] == scenario_id
        )
        assert entity["states"]["waived"]["satisfied"] is True
        assert entity["states"]["implemented"]["satisfied"] is False
        assert entity["states"]["tested"]["satisfied"] is False
        assert entity["states"]["accepted"]["satisfied"] is False
        revoke_payload = {
            "project_id": "p1", "expected_revision": change["revision"],
            "client_operation_id": "5" * 32,
            "reviewer_label": "reviewer", "reason": "No longer needed",
        }
        revoked = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/coverage/waivers/"
            f"{waiver['waiver_id']}/revoke",
            json=revoke_payload,
        )
        assert revoked.status == 200, await revoked.text()
        assert (await revoked.json())["state"] == "REVOKED"
        replay_revoked = await client.post(
            f"/api/v1/pm/changes/{change['change_id']}/coverage/waivers/"
            f"{waiver['waiver_id']}/revoke",
            json=revoke_payload,
        )
        assert replay_revoked.status == 200
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_change_get_projects_real_approval_and_invalidates_after_edit(tmp_path: Path, monkeypatch) -> None:
    from symphony.pm.models import ValidationReport

    monkeypatch.setattr(routes_module, "_openspec_readiness", lambda *_args: ValidationReport(True))
    client, planning = await _authoring_client(tmp_path)
    registry = RunRegistry(tmp_path / "project/.symphony/state.db")
    repository = PmDocumentRepository(project_id="p1", project_root=tmp_path / "project",
                                     registry=registry, planning_workspace_root=planning.parent)
    change_id = uuid.uuid4().hex
    repository.claim_workspace(change_id, planning, "owner")
    change = _populate(repository, change_id, planning)
    path = f"/api/v1/pm/changes/{change_id}"
    headers = {"X-Symphony-PM-Owner": "owner"}
    envelope = {"project_id": "p1", "change_id": change_id, "expected_revision": change.revision}
    try:
        initial = await (await client.get(path + "?project_id=p1", headers=headers)).json()
        approval = await client.post(path + "/approvals", headers=headers, json={
            **envelope, "confirmation_token": initial["approval_confirmation_token"], "reviewer_label": "reviewer",
        })
        assert approval.status == 201, await approval.text()
        manifest = await approval.json()
        response = await client.get(path + "?project_id=p1", headers=headers)
        assert response.status == 200, await response.text()
        approved = await response.json()
        assert (approved["state"], approved["approval_id"]) == ("Approved", manifest["approval_id"])
        assert (approved["scope_hash"], approved["revision"]) == (manifest["scope_hash"], manifest["revision"])
        source = repository.get_document(PmRequestContext("p1", change_id), "proposal", owner="owner").source
        patched = await client.patch(path + "/documents/proposal", headers=headers, json={
            **envelope, "client_operation_id": "approval-api-edit",
            "payload": {"operations": [{"op": "replace_source", "source": source + "\nReview note: edited.\n"}]},
        })
        assert patched.status == 200, await patched.text()
        current_response = await client.get(path + "?project_id=p1", headers=headers)
        assert current_response.status == 200, await current_response.text()
        current = await current_response.json()
        assert current["revision"] != approved["revision"]
        assert (current["state"], current["approval_id"]) == ("Validated", None)
    finally:
        registry.close()
        await client.close()


@pytest.mark.timing
@pytest.mark.asyncio
async def test_R32_04_http_writer_does_not_block_loop_waiting_for_snapshot_lock(tmp_path):
    client, workspace = await _authoring_client(tmp_path)
    registry = RunRegistry(tmp_path / "project/.symphony/state.db")
    repo = PmDocumentRepository(project_id="p1", project_root=tmp_path / "project", registry=registry, planning_workspace_root=tmp_path / "planning-root")
    cid = "b" * 32
    repo.claim_workspace(cid, workspace, "local-operator")
    change = repo.create_change(change_id=cid, locator="lock-test", title="Lock test", owner="local-operator")
    root = workspace / "openspec/changes/lock-test"
    ready = threading.Event()
    def reader():
        with __import__("symphony.pm.documents", fromlist=["_change_lock"])._change_lock(root, shared=True):
            ready.set()
            time.sleep(.15)
    thread = threading.Thread(target=reader)
    thread.start()
    assert ready.wait(1)
    delays = []
    async def heartbeat():
        previous = time.monotonic()
        while True:
            await asyncio.sleep(.01)
            now = time.monotonic()
            delays.append(now - previous)
            previous = now
    beat = asyncio.create_task(heartbeat())
    try:
        response = await client.patch(f"/api/v1/pm/changes/{cid}/documents/proposal", json={"project_id":"p1", "expected_revision":change.revision, "client_operation_id":"lock-patch", "payload":{"operations":[{"op":"replace_source", "source":"# Lock test\n\n## Problem\nPreserve source.\n"}]}})
        payload = await response.text()
        await asyncio.sleep(.02)
        print("HTTP_WRITER", response.status, payload[:160], "MAX_HEARTBEAT_SECONDS", max(delays))
        assert response.status == 200
        assert max(delays) < .07, "HTTP patch waits on a shared snapshot lock on the event loop"
    finally:
        beat.cancel()
        await asyncio.gather(beat, return_exceptions=True)
        thread.join(1)
        registry.close()
        await client.close()


@pytest.mark.timing
@pytest.mark.asyncio
async def test_N6_report_git_is_bounded_and_off_event_loop(tmp_path, monkeypatch):
    client = await _client(tmp_path)
    blob = b'{"result":"passed"}\n'
    row = {"project_id": "p1", "tested_commit": "a" * 40, "report_ref_json": json.dumps({"ref_kind": "git_blob", "ref": "manual-report.json", "sha256": "sha256:" + hashlib.sha256(blob).hexdigest()})}
    monkeypatch.setattr(RunRegistry, "get_pm_manual_verification_receipt", lambda *_args: row)
    original = subprocess.run
    observed = []
    def slow_git(argv, *args, **kwargs):
        if argv[0] == "git" and "show" in argv:
            observed.append({"thread": threading.current_thread().name, "timeout": kwargs.get("timeout")})
            time.sleep(.15)
            return subprocess.CompletedProcess(argv, 0, stdout=blob, stderr=b"")
        return original(argv, *args, **kwargs)
    monkeypatch.setattr(routes_module.subprocess, "run", slow_git)
    delays = []
    async def heartbeat():
        previous = time.monotonic()
        while True:
            await asyncio.sleep(.01)
            now = time.monotonic()
            delays.append(now - previous)
            previous = now
    beat = asyncio.create_task(heartbeat())
    try:
        response = await client.get("/api/v1/pm/evidence/manual-receipt/report?project_id=p1")
        assert response.status == 200 and await response.read() == blob
        await asyncio.sleep(.02)
        print("REPORT_GIT", observed, "MAX_HEARTBEAT_SECONDS", max(delays))
        assert observed[0]["timeout"] is not None and observed[0]["thread"] != "MainThread" and max(delays) < .07
    finally:
        beat.cancel()
        await asyncio.gather(beat, return_exceptions=True)
        await client.close()



@pytest.mark.asyncio
async def test_concurrent_http_writers_keep_one_revision_and_refuse_stale_retry(tmp_path):
    client, workspace = await _authoring_client(tmp_path)
    registry = RunRegistry(tmp_path / "project/.symphony/state.db")
    repository = PmDocumentRepository(project_id="p1", project_root=tmp_path / "project", registry=registry, planning_workspace_root=tmp_path / "planning-root")
    change_id = "c" * 32
    repository.claim_workspace(change_id, workspace, "local-operator")
    change = repository.create_change(change_id=change_id, locator="concurrent-writers", title="Writers", owner="local-operator")
    url = f"/api/v1/pm/changes/{change_id}/documents/proposal"
    bodies = [{"project_id": "p1", "expected_revision": change.revision, "client_operation_id": f"writer-{index}", "payload": {"operations": [{"op": "replace_source", "source": f"# Writer {index}\n\n## Problem\nSave exactly one revision.\n"}]}} for index in range(2)]
    try:
        responses = await asyncio.gather(*(client.patch(url, json=body) for body in bodies))
        assert sorted(response.status for response in responses) == [200, 409]
        winner = next(index for index, response in enumerate(responses) if response.status == 200)
        loser = 1 - winner
        refusal = await responses[loser].json()
        assert refusal["error"]["code"] == "revision_conflict"
        saved = await client.get(url + "?project_id=p1")
        assert (await saved.json())["source"] == bodies[winner]["payload"]["operations"][0]["source"]
        retry = await client.patch(url, json=bodies[loser])
        assert retry.status == 409 and (await retry.json())["error"]["code"] == "revision_conflict"
        current = await client.get(f"/api/v1/pm/changes/{change_id}?project_id=p1")
        recovered = await current.json()
        bodies[loser]["expected_revision"] = recovered["revision"]
        corrected = await client.patch(url, json=bodies[loser])
        assert corrected.status == 200
    finally:
        registry.close()
        await client.close()


@pytest.mark.asyncio
async def test_report_git_timeout_and_missing_executable_preserve_refusal_and_retry(tmp_path, monkeypatch):
    client = await _client(tmp_path)
    blob = b"immutable manual report\n"
    row = {"project_id": "p1", "tested_commit": "a" * 40, "report_ref_json": json.dumps({"ref_kind": "git_blob", "ref": "manual-report.txt", "sha256": "sha256:" + hashlib.sha256(blob).hexdigest()})}
    monkeypatch.setattr(RunRegistry, "get_pm_manual_verification_receipt", lambda *_args: row)
    original = subprocess.run
    def unavailable(argv, *args, **kwargs):
        if argv[0] == "git" and "show" in argv:
            assert kwargs.get("timeout") == 30
            raise subprocess.TimeoutExpired(argv, kwargs["timeout"])
        return original(argv, *args, **kwargs)
    try:
        monkeypatch.setattr(routes_module.subprocess, "run", unavailable)
        response = await client.get("/api/v1/pm/evidence/manual-receipt/report?project_id=p1")
        assert response.status == 404 and (await response.json())["error"]["code"] == "evidence_not_found"
        def missing(argv, *args, **kwargs):
            if argv[0] == "git" and "show" in argv:
                raise FileNotFoundError("git unavailable")
            return original(argv, *args, **kwargs)
        monkeypatch.setattr(routes_module.subprocess, "run", missing)
        response = await client.get("/api/v1/pm/evidence/manual-receipt/report?project_id=p1")
        assert response.status == 404 and (await response.json())["error"]["code"] == "evidence_not_found"
        def recovered(argv, *args, **kwargs):
            if argv[0] == "git" and "show" in argv:
                return subprocess.CompletedProcess(argv, 0, stdout=blob, stderr=b"")
            return original(argv, *args, **kwargs)
        monkeypatch.setattr(routes_module.subprocess, "run", recovered)
        response = await client.get("/api/v1/pm/evidence/manual-receipt/report?project_id=p1")
        assert response.status == 200 and await response.read() == blob
        row["project_id"] = "other"
        assert (await client.get("/api/v1/pm/evidence/manual-receipt/report?project_id=p1")).status == 409
    finally:
        await client.close()


@pytest.mark.timing
@pytest.mark.asyncio
async def test_manual_preview_confirm_run_git_off_loop_and_keep_actual_receipt_binding(tmp_path, monkeypatch):
    client, registry, _planning, path, headers, ticket = await _published_export_fixture(tmp_path, monkeypatch)
    mapping = registry.get_pm_managed_ticket("p1", ticket)
    assert mapping is not None
    change = await (await client.get(path + "?project_id=p1", headers=headers)).json()
    project = tmp_path / "project"
    commit = subprocess.check_output(["git", "-C", str(project), "rev-parse", "HEAD"], text=True).strip()
    definition = {"procedure_id": "manual-ux", "kind": "manual", "instructions": "Review immutable fixture report", "report": "required", "artifacts": "optional"}
    definition["definition_digest"] = "sha256:" + hashlib.sha256(json.dumps(definition, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    registry._connect().execute("UPDATE pm_handoff_items SET required_procedures_json=? WHERE ticket_identifier=?", (json.dumps({"manual-ux": definition}), ticket))
    payload = {"project_id": "p1", "expected_revision": change["revision"], "ticket_identifier": ticket, "task_revision": mapping.task_revision, "tested_commit": commit, "procedure_id": "manual-ux", "environment_summary": "disposable test", "started_at": "2026-10-14T00:00:00Z", "completed_at": "2026-10-14T00:01:00Z", "result": "passed", "report_ref": {"ref_kind": "git_blob", "ref": "README.md"}, "artifact_refs": []}
    original = subprocess.run
    observed = []
    def slow_git(argv, *args, **kwargs):
        if argv[0] == "git" and "show" in argv:
            observed.append((threading.current_thread().name, kwargs.get("timeout")))
            time.sleep(.12)
        return original(argv, *args, **kwargs)
    monkeypatch.setattr(routes_module.subprocess, "run", slow_git)
    delays = []
    async def heartbeat():
        previous = time.monotonic()
        while True:
            await asyncio.sleep(.01)
            now = time.monotonic()
            delays.append(now - previous)
            previous = now
    beat = asyncio.create_task(heartbeat())
    try:
        preview_response = await client.post(path + "/manual-evidence/preview", headers=headers, json=payload)
        assert preview_response.status == 200
        preview = await preview_response.json()
        assert preview["ready"] is True, preview
        confirmed = await client.post(path + "/manual-evidence", headers=headers, json={**payload, "confirmation_token": preview["confirmation_token"], "client_operation_id": "manual-thread-receipt"})
        assert confirmed.status == 201, await confirmed.text()
        receipt = await confirmed.json()
        assert receipt["tested_commit"] == commit and receipt["task_revision"] == mapping.task_revision
        assert receipt["return_code"] is None and receipt["run_id"] is None
        assert observed and all(name != "MainThread" and timeout == 30 for name, timeout in observed)
        await asyncio.sleep(.02)
        assert max(delays) < .07, delays
    finally:
        beat.cancel()
        await asyncio.gather(beat, return_exceptions=True)
        registry.close()
        await client.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("disconnect", [False, True])
async def test_cancel_waits_for_real_admission_and_disconnected_request(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, disconnect: bool,
) -> None:
    from symphony.pm.planning import PmPlanningClient
    from tests.test_pm_planning_cancellation import MemoryRuntime

    runtime = MemoryRuntime("admission")
    runtime.install(monkeypatch)
    auth = tmp_path / "dummy-auth"
    auth.write_text("{}")
    planning_client = PmPlanningClient(auth_file=auth, timeout_seconds=240)
    monkeypatch.setattr(routes_module, "PmPlanningClient", lambda **_kwargs: planning_client)
    cancel_entered = asyncio.Event()
    original_operation = routes_module._operation
    operation_id = "9" * 32

    def observe_operation(registry, requested_id, **kwargs):
        row = original_operation(registry, requested_id, **kwargs)
        if requested_id == operation_id:
            cancel_entered.set()
        return row

    monkeypatch.setattr(routes_module, "_operation", observe_operation)
    client, workspace = await _authoring_client(tmp_path)
    proposal_request = None
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(workspace),
            "locator": "cancel-before-admission", "title": "Cancel before admission",
        })
        assert created.status == 201, await created.text()
        change = await created.json()
        proposal_request = asyncio.create_task(client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
            json={"project_id": "p1", "change_id": change["change_id"],
                  "expected_revision": change["revision"], "client_operation_id": operation_id,
                  "action": "find_gaps", "document_id": "proposal",
                  "impacted_document_ids": ["proposal"], "instruction": "Review"},
        ))
        await asyncio.wait_for(runtime.entered.wait(), 2)
        if disconnect:
            proposal_request.cancel()
            await asyncio.gather(proposal_request, return_exceptions=True)
        cancel = asyncio.create_task(client.post(f"/api/v1/pm/operations/{operation_id}/cancel"))
        await asyncio.wait_for(cancel_entered.wait(), 2)
        await asyncio.sleep(0)
        runtime.release.set()
        response = await asyncio.wait_for(cancel, 2)
        assert response.status == 200, await response.text()
        assert (await response.json())["state"] == "cancelled"
        if not disconnect:
            proposal = await asyncio.wait_for(proposal_request, 2)
            assert proposal.status == 409, await proposal.text()
            assert (await proposal.json())["error"]["code"] == "operation_cancelled"
        assert runtime.spawns == 0, "cancel acknowledged before startup was fenced"
        assert not runtime.admission_cancelled
        repeated = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert repeated.status == 200
        assert (await repeated.json())["state"] == "cancelled"
        assert planning_client.process is None and not planning_client.cleanup_uncertain
    finally:
        runtime.release.set()
        if proposal_request is not None:
            await asyncio.gather(proposal_request, return_exceptions=True)
        await runtime.close()
        await client.close()


@pytest.mark.asyncio
async def test_cancel_cleanup_retry_clears_retained_error_and_allows_new_operation(tmp_path, monkeypatch):
    from symphony.pm.planning import PmPlanningClient
    from tests.test_pm_planning_cancellation import MemoryRuntime

    runtime = MemoryRuntime("turn_wait")
    runtime.install(monkeypatch)
    auth = tmp_path / "dummy-auth"
    auth.write_text("{}")
    first_client = PmPlanningClient(auth_file=auth, timeout_seconds=240)
    second_client = PmPlanningClient(auth_file=auth, timeout_seconds=240)
    clients = [first_client, second_client]
    monkeypatch.setattr(routes_module, "PmPlanningClient", lambda **_kwargs: clients.pop(0))
    client, workspace = await _authoring_client(tmp_path)
    registry = RunRegistry(tmp_path / "project/.symphony/state.db")
    next_runtime = None
    proposal = None
    operation_id = "7" * 32
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(workspace),
            "locator": "cancel-cleanup-retry", "title": "Cancel cleanup retry",
        })
        change = await created.json()
        payload = {"project_id": "p1", "change_id": change["change_id"],
                   "expected_revision": change["revision"], "client_operation_id": operation_id,
                   "action": "find_gaps", "document_id": "proposal",
                   "impacted_document_ids": ["proposal"], "instruction": "Review"}
        path = f"/api/v1/pm/changes/{change['change_id']}/planning-actions"
        proposal = asyncio.create_task(client.post(path, json=payload))
        await runtime.entered.wait()
        import symphony.pm.planning as planning_module
        monkeypatch.setattr(planning_module, "process_group_exists", lambda _pid: None)
        failed = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert failed.status == 503
        assert (await failed.json())["error"]["code"] == "planning_cleanup_failed"
        original = await proposal
        assert original.status == 503
        row = routes_module._operation(registry, operation_id, project_id="p1")
        assert row["state"] == "pending"
        blocked = await client.post(path, json={**payload, "client_operation_id": "8" * 32})
        assert blocked.status == 503
        assert (await blocked.json())["error"]["code"] == "planning_cleanup_failed"
        assert len(clients) == 1
        monkeypatch.setattr(planning_module, "process_group_exists", lambda _pid: False)
        retried = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert retried.status == 200, await retried.text()
        assert (await retried.json())["state"] == "cancelled"
        assert first_client.process is None and not first_client.cleanup_uncertain
        next_runtime = MemoryRuntime("none")
        next_runtime.patch = ""
        next_runtime.install(monkeypatch)
        accepted = await client.post(path, json={**payload, "client_operation_id": "8" * 32})
        assert accepted.status == 201, await accepted.text()
        assert (await accepted.json())["state"] == "completed"
    finally:
        runtime.release.set()
        if proposal is not None:
            await asyncio.gather(proposal, return_exceptions=True)
        await runtime.close()
        if next_runtime is not None:
            await next_runtime.close()
        registry.close()
        await client.close()


@pytest.mark.asyncio
async def test_cancel_finishes_original_http_request_thread_within_unchanged_join(tmp_path, monkeypatch):
    import threading
    import urllib.error
    import urllib.request

    from symphony.pm.planning import PmPlanningClient
    from tests.test_pm_planning_cancellation import MemoryRuntime

    runtime = MemoryRuntime("admission")
    runtime.install(monkeypatch)
    auth = tmp_path / "dummy-auth"
    auth.write_text("{}")
    planning_client = PmPlanningClient(auth_file=auth, timeout_seconds=240)
    monkeypatch.setattr(routes_module, "PmPlanningClient", lambda **_kwargs: planning_client)
    cancel_entered = asyncio.Event()
    original_operation = routes_module._operation
    operation_id = "6" * 32

    def observe_operation(registry, requested_id, **kwargs):
        row = original_operation(registry, requested_id, **kwargs)
        if requested_id == operation_id:
            cancel_entered.set()
        return row

    monkeypatch.setattr(routes_module, "_operation", observe_operation)
    client, workspace = await _authoring_client(tmp_path)
    result = {}
    thread = None
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(workspace),
            "locator": "cancel-http-thread", "title": "Cancel HTTP thread",
        })
        change = await created.json()
        url = str(client.make_url(f"/api/v1/pm/changes/{change['change_id']}/planning-actions"))
        body = json.dumps({"project_id": "p1", "change_id": change["change_id"],
                           "expected_revision": change["revision"], "client_operation_id": operation_id,
                           "action": "find_gaps", "document_id": "proposal",
                           "impacted_document_ids": ["proposal"], "instruction": "Review"}).encode()

        def post_proposal():
            try:
                request = urllib.request.Request(url, data=body, headers={"Content-Type": "application/json"})
                with urllib.request.urlopen(request, timeout=240) as response:
                    result["status"] = response.status
            except urllib.error.HTTPError as error:
                result["status"] = error.code
                result["error"] = json.loads(error.read())["error"]["code"]
            except Exception as error:
                result["exception"] = repr(error)

        thread = threading.Thread(target=post_proposal, daemon=True)
        thread.start()
        await asyncio.wait_for(runtime.entered.wait(), 2)
        cancelled = asyncio.create_task(client.post(f"/api/v1/pm/operations/{operation_id}/cancel"))
        await asyncio.wait_for(cancel_entered.wait(), 2)
        await asyncio.sleep(0)
        runtime.release.set()
        response = await asyncio.wait_for(cancelled, 2)
        assert response.status == 200, await response.text()
        assert (await response.json())["state"] == "cancelled"
        await asyncio.to_thread(thread.join, 30)
        assert not thread.is_alive(), "cancelled planning request did not terminate"
        assert result == {"status": 409, "error": "operation_cancelled"}
        assert runtime.spawns == 0
    finally:
        runtime.release.set()
        if thread is not None:
            await asyncio.to_thread(thread.join, 30)
            assert not thread.is_alive()
        await runtime.close()
        await client.close()


@pytest.mark.asyncio
async def test_cancel_refuses_handler_cleanup_failure_then_explicit_retry_recovers(tmp_path, monkeypatch):
    from symphony.pm.models import PmError
    from symphony.pm.planning import PmPlanningClient
    from tests.test_pm_planning_cancellation import MemoryRuntime

    class HandlerCleanupClient(PmPlanningClient):
        handler = None

        async def propose(self, *args, **kwargs):
            self.handler = asyncio.current_task()
            return await super().propose(*args, **kwargs)

        async def stop(self):
            if asyncio.current_task() is self.handler:
                raise PmError("planning_cleanup_failed", "handler cleanup uncertain", status=503)
            await super().stop()

    runtime = MemoryRuntime("turn_wait")
    runtime.install(monkeypatch)
    auth = tmp_path / "dummy-auth"
    auth.write_text("{}")
    planning_client = HandlerCleanupClient(auth_file=auth, timeout_seconds=240)
    monkeypatch.setattr(routes_module, "PmPlanningClient", lambda **_kwargs: planning_client)
    client, workspace = await _authoring_client(tmp_path)
    proposal = None
    operation_id = "5" * 32
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(workspace),
            "locator": "handler-cleanup", "title": "Handler cleanup",
        })
        change = await created.json()
        proposal = asyncio.create_task(client.post(
            f"/api/v1/pm/changes/{change['change_id']}/planning-actions",
            json={"project_id": "p1", "change_id": change["change_id"],
                  "expected_revision": change["revision"], "client_operation_id": operation_id,
                  "action": "find_gaps", "document_id": "proposal", "impacted_document_ids": ["proposal"]},
        ))
        await runtime.entered.wait()
        failed = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert failed.status == 503
        assert (await failed.json())["error"]["code"] == "planning_cleanup_failed"
        assert (await proposal).status == 503
        state = await client.get(f"/api/v1/pm/operations/{operation_id}")
        assert (await state.json())["state"] == "pending"
        retried = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert retried.status == 200, await retried.text()
        assert (await retried.json())["state"] == "cancelled"
        assert planning_client.process is None and not planning_client.cleanup_uncertain
    finally:
        runtime.release.set()
        if proposal is not None:
            await asyncio.gather(proposal, return_exceptions=True)
        await runtime.close()
        await client.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("stage", ["admission", "spawn"])
@pytest.mark.parametrize("repeated_caller_cancel", [False, True])
async def test_disconnected_caller_deadline_retains_unfinished_startup_until_explicit_retry(
    tmp_path, monkeypatch, stage, repeated_caller_cancel,
):
    from symphony.pm.models import PmError
    from symphony.pm.planning import PmPlanningClient
    from tests.test_pm_planning_cancellation import MemoryRuntime

    class ObservedClient(PmPlanningClient):
        handler = None

        async def propose(self, *args, **kwargs):
            self.handler = asyncio.current_task()
            return await super().propose(*args, **kwargs)

    class UnexpectedSecondClient(PmPlanningClient):
        async def propose(self, *args, **kwargs):
            raise PmError("planning_failed", "a second client was admitted", status=503)

    runtime = MemoryRuntime(stage)
    runtime.install(monkeypatch)
    auth = tmp_path / "dummy-auth"
    auth.write_text("{}")
    first = ObservedClient(auth_file=auth, timeout_seconds=0.02)
    constructed = []
    recovered = False

    def factory(**_kwargs):
        if not constructed:
            created = first
        elif recovered:
            created = PmPlanningClient(auth_file=auth, timeout_seconds=240)
        else:
            created = UnexpectedSecondClient(auth_file=auth)
        constructed.append(created)
        return created

    monkeypatch.setattr(routes_module, "PmPlanningClient", factory)
    client, workspace = await _authoring_client(tmp_path)
    proposal = None
    next_runtime = None
    operation_id = "4" * 32
    try:
        created = await client.post("/api/v1/pm/changes", json={
            "project_id": "p1", "workspace_path": str(workspace),
            "locator": "caller-deadline", "title": "Caller deadline",
        })
        change = await created.json()
        payload = {"project_id": "p1", "change_id": change["change_id"],
                   "expected_revision": change["revision"], "client_operation_id": operation_id,
                   "action": "find_gaps", "document_id": "proposal",
                   "impacted_document_ids": ["proposal"]}
        path = f"/api/v1/pm/changes/{change['change_id']}/planning-actions"
        proposal = asyncio.create_task(client.post(path, json=payload))
        await asyncio.wait_for(runtime.entered.wait(), 2)
        proposal.cancel()
        await asyncio.gather(proposal, return_exceptions=True)
        await asyncio.wait_for(first._cancel_requested.wait(), 1)
        if repeated_caller_cancel:
            first.handler.cancel()
        await asyncio.wait_for(asyncio.gather(asyncio.shield(first.handler), return_exceptions=True), 1)
        assert first.process is None and first.cleanup_uncertain
        assert first._proposal_task is not None and not first._proposal_task.done()
        if stage == "spawn":
            assert first._temp is not None and Path(first._temp.name).is_dir()
        state = await client.get(f"/api/v1/pm/operations/{operation_id}")
        actual_state = (await state.json())["state"]
        blocked = await client.post(path, json={**payload, "client_operation_id": "3" * 32})
        blocked_body = await blocked.json()
        assert (actual_state, blocked.status, blocked_body["error"]["code"], len(constructed)) == (
            "pending", 503, "planning_cleanup_failed", 1,
        )
        refused = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert refused.status == 503
        assert first.cleanup_uncertain and len(constructed) == 1
        runtime.release.set()
        await asyncio.wait_for(asyncio.gather(first._proposal_task, return_exceptions=True), 1)
        retried = await client.post(f"/api/v1/pm/operations/{operation_id}/cancel")
        assert retried.status == 200, await retried.text()
        assert (await retried.json())["state"] == "cancelled"
        assert not first.cleanup_uncertain and first.process is None
        assert "turn/start" not in runtime.methods and not runtime.admission_cancelled
        recovered = True
        next_runtime = MemoryRuntime("none")
        next_runtime.patch = ""
        next_runtime.install(monkeypatch)
        accepted = await client.post(path, json={**payload, "client_operation_id": "2" * 32})
        assert accepted.status == 201, await accepted.text()
        assert len(constructed) == 2
    finally:
        runtime.release.set()
        if proposal is not None:
            await asyncio.gather(proposal, return_exceptions=True)
        if first._proposal_task is not None:
            await asyncio.gather(first._proposal_task, return_exceptions=True)
        await first.cancel_proposal()
        await runtime.close()
        if next_runtime is not None:
            await next_runtime.close()
        await client.close()
