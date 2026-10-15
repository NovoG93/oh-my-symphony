from __future__ import annotations

import io
import re
import zipfile
from dataclasses import replace
from datetime import UTC, datetime

import openpyxl
import pytest

from symphony.pm.exchange import MAX_UNCOMPRESSED, WorkbookExchange
from symphony.pm.models import ChangeView, DocumentView, PmError
from symphony.pm.routes import _render_export


def _rewrite_member(data: bytes, name: str, transform) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info)
            target.writestr(info, transform(payload) if info.filename == name else payload)
    return output.getvalue()


def _change(title: str = "Implement API", revision: str = "sha256:base") -> ChangeView:
    fence = "`" * 3
    source = f"""# Plan
<!-- symphony-wbs-format: 1 -->
<!-- symphony-group:G-ABCDEFGH -->
## 1. Group
{fence}yaml symphony-group
id: G-ABCDEFGH
parent_id: null
order: 10
{fence}
<!-- /symphony-group -->
<!-- symphony-task:T-ABCDEFGH -->
- [ ] {title}
{fence}yaml symphony-task
id: T-ABCDEFGH
parent_id: G-ABCDEFGH
order: 10
type: Backend
depends_on:
  - T-ZZZZZZZZ
requirements:
  - REQ-1
{fence}
### Scope
Base description.
### Acceptance criteria
- Base acceptance.
### Verification
- Base verification.
### Integration impact
Base integration.
### Risk
Base risk.
### Notes
Base notes.
<!-- /symphony-task -->
<!-- symphony-task:T-ZZZZZZZZ -->
- [ ] Foundation
{fence}yaml symphony-task
id: T-ZZZZZZZZ
parent_id: G-ABCDEFGH
order: 20
type: Backend
depends_on: []
{fence}
### Scope
Foundation scope.
### Acceptance criteria
- Foundation works.
### Verification
- Run foundation tests.
### Integration impact
None.
<!-- /symphony-task -->
"""
    return ChangeView(
        project_id="p1",
        change_id="c1",
        locator="demo",
        revision=revision,
        scope_hash="sha256:scope",
        state="Draft",
        base_git_sha="a" * 40,
        schema_hash="schema",
        documents=(
            DocumentView("tasks", "tasks.md", source, "hash", "draft", True),
        ),
    )


def test_six_sheet_export_forces_formula_prefix_to_literal_and_round_trips() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change("=SUM(1,1)"), {})
    workbook = openpyxl.load_workbook(
        io.BytesIO(payload.data), read_only=True, data_only=False, keep_links=False
    )
    try:
        assert tuple(workbook.sheetnames) == (
            "Read Me",
            "WBS",
            "Task Details",
            "Dependencies",
            "Traceability",
            "Metadata",
        )
        title = workbook["WBS"]["F3"]
        assert title.value == "=SUM(1,1)"
        assert title.data_type == "s"
    finally:
        workbook.close()

    candidate = exchange.preview_import(payload.data, _change(), _change())
    assert candidate.base_revision == "sha256:base"
    assert candidate.filename == "workbook.xlsx"
    assert candidate.conflicts == ()

    named = exchange.preview_import(
        payload.data, _change(), _change(), filename="reviewed.XLSX"
    )
    assert named.filename == "reviewed.XLSX"
    assert named.preview_digest != candidate.preview_digest


def test_export_id_is_embedded_for_durable_trusted_base_lookup() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change(), {}, export_id="a" * 32)
    metadata = exchange.inspect_metadata(payload.data)
    assert metadata["export_id"] == "a" * 32
    assert metadata["revision"] == "sha256:base"


def test_export_readme_counts_only_leaf_runtime_snapshots_and_records_time() -> None:
    before = datetime.now(UTC)
    payload = WorkbookExchange().export(_change(), {
        "G-ABCDEFGH": {"state": "Done"},
        "T-ABCDEFGH": {"state": "=Custom state"},
        "unknown-task": {"state": "Done"},
    })
    after = datetime.now(UTC)
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    try:
        rows = dict(workbook["Read Me"].values)
        assert before <= datetime.fromisoformat(rows["Exported at (UTC)"]) <= after
        assert rows["Leaf tasks"] == 2
        assert rows["Not supplied"] == 1
        assert rows["Supplied: =Custom state"] == 1
        assert "Supplied: Done" not in rows
        assert "Supplied: Draft" not in rows
        assert "export-time snapshots" in rows["Execution snapshots"]
        assert "not supplied" in rows["Status legend"].lower()
        assert workbook["WBS"]["L3"].value == "=Custom state"
        assert all(cell.data_type != "f" for row in workbook["Read Me"] for cell in row)
    finally:
        workbook.close()


@pytest.mark.parametrize("runtime", [{}, {"T-ABCDEFGH": {"state": ""}}, {"T-ABCDEFGH": {"state": None}}])
def test_export_missing_runtime_is_not_inferred_from_planning_state(runtime) -> None:
    change = replace(_change(), state="Approved")
    payload = WorkbookExchange().export(change, runtime)
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    try:
        rows = dict(workbook["Read Me"].values)
        assert rows["State"] == "Approved"
        assert rows["Not supplied"] == rows["Leaf tasks"] == 2
        assert workbook["WBS"]["L3"].value is None
    finally:
        workbook.close()


def test_export_empty_and_many_distinct_statuses_stay_within_readme_bounds() -> None:
    exchange = WorkbookExchange()
    empty = replace(_change(), documents=())
    workbook = openpyxl.load_workbook(io.BytesIO(exchange.export(empty, {}).data))
    try:
        assert dict(workbook["Read Me"].values)["Leaf tasks"] == 0
        assert workbook["Read Me"].max_row <= 64
    finally:
        workbook.close()

    change = _change()
    document = change.documents[0]
    group = document.source.split("<!-- symphony-task:", 1)[0]
    foundation = document.source.split("<!-- symphony-task:T-ZZZZZZZZ -->", 1)[1]
    ids = [f"T-{index:08d}" for index in range(80)]
    source = group + "".join(
        f"<!-- symphony-task:{task_id} -->" + foundation.replace("T-ZZZZZZZZ", task_id)
        for task_id in ids
    )
    change = replace(change, documents=(replace(document, source=source),))
    runtime = {task_id: {"state": f"@state {index} " + "x" * 200} for index, task_id in enumerate(ids)}
    payload = exchange.export(change, runtime)
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    try:
        readme = workbook["Read Me"]
        rows = dict(readme.values)
        assert readme.max_row <= 64 and readme.max_column == 2
        assert rows["Leaf tasks"] == 80
        assert rows["Additional supplied statuses"] > 0
        summary_start = next(row[0].row for row in readme if row[0].value == "Leaf status summary")
        assert sum(row[1].value for row in list(readme)[summary_start:]) == 80
        assert workbook["WBS"]["L3"].value == runtime[ids[0]]["state"]
        assert all(cell.data_type != "f" for row in readme for cell in row)
    finally:
        workbook.close()
    assert exchange.preview_import(payload.data, change, change).changes == ()


def test_export_native_presentation_preserves_headers_and_advisory_controls() -> None:
    workbook = openpyxl.load_workbook(io.BytesIO(WorkbookExchange().export(_change(), {}).data))
    try:
        assert {name: workbook[name].freeze_panes for name in ("WBS", "Task Details", "Dependencies", "Traceability")} == {
            "WBS": "D2", "Task Details": "B2", "Dependencies": "C2", "Traceability": "B2",
        }
        assert workbook["WBS"]["F2"].alignment.indent == 0
        assert workbook["WBS"]["F3"].alignment.indent == 1
        assert workbook["WBS"].row_dimensions[3].outlineLevel == 1
        assert workbook["WBS"]["B3"].fill.fgColor != workbook["WBS"]["F3"].fill.fgColor
        assert workbook["WBS"]["L3"].fill.fgColor == workbook["WBS"]["B3"].fill.fgColor
        assert workbook["Traceability"]["D2"].fill.fgColor != workbook["Traceability"]["C2"].fill.fgColor
        assert workbook["Metadata"]["A2"].fill.fgColor == workbook["WBS"]["B3"].fill.fgColor
        assert workbook["WBS"]["F3"].alignment.wrap_text
        assert workbook["WBS"].auto_filter.ref == "A1:N4"
        expected = {
            "WBS": {"A": '"Keep,Add,Delete"'},
            "Dependencies": {"C": '"requires_code,order_only"', "E": '"Keep,Add,Delete"'},
            "Traceability": {"B": '"requirement,scenario,storyboard,design,api"', "E": '"Keep,Add,Delete"'},
        }
        for name, columns in expected.items():
            validations = workbook[name].data_validations.dataValidation
            assert len(validations) == len(columns)
            for validation in validations:
                column = str(validation.sqref).split("2:")[0]
                assert validation.type == "list" and validation.formula1 == columns[column]
                assert validation.allowBlank and validation.showInputMessage
                assert not validation.showErrorMessage and not validation.showDropDown
        assert tuple(cell.value for cell in workbook["Metadata"][1]) == (
            "format_version", "project_id", "change_id", "revision", "scope_hash", "export_id",
        )
        assert workbook["Metadata"].max_row == 2 and workbook["Metadata"].max_column == 6
    finally:
        workbook.close()


@pytest.mark.parametrize("legacy", [False, True])
@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_export_custom_values_literal_prefixes_and_legacy_v1_still_round_trip(legacy, prefix) -> None:
    exchange = WorkbookExchange()
    change = _change(prefix + "Custom title")
    document = change.documents[0]
    source = document.source.replace("type: Backend", "type: CustomDelivery").replace(
        "type: CustomDelivery", "type: CustomDelivery\npriority: urgent-custom\nagent_kind: bespoke-agent", 1,
    )
    change = replace(change, documents=(replace(document, source=source),))
    payload = exchange.export(change, {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    try:
        assert workbook["WBS"]["E3"].value == "CustomDelivery"
        assert workbook["WBS"]["I3"].value == "urgent-custom"
        assert workbook["WBS"]["H3"].value == "bespoke-agent"
        assert workbook["WBS"]["F3"].data_type == "s"
        workbook["WBS"]["A3"] = None
        workbook["Dependencies"]["C2"] = None
        workbook["Dependencies"]["E2"] = None
        workbook["Traceability"]["B2"] = " ReQuIrEmEnT "
        workbook["Traceability"]["E2"] = None
        if legacy:
            workbook["Read Me"].delete_rows(6, 64)
            workbook["WBS"].delete_cols(13, 2)
            workbook["Dependencies"].delete_cols(6, 2)
            for sheet in workbook:
                sheet.freeze_panes = "A2"
                sheet.data_validations.dataValidation.clear()
        output = io.BytesIO()
        workbook.save(output)
    finally:
        workbook.close()
    candidate = exchange.preview_import(output.getvalue(), change, change)
    assert candidate.changes == () and candidate.conflicts == () and candidate.findings == ()


def test_export_derived_fields_are_snapshot_only_and_ignored_on_import() -> None:
    exchange = WorkbookExchange()
    change = _change("=SUM(1,1)")
    payload = exchange.export(change, {"T-ABCDEFGH": {"state": "Done", "link": "/api/v1/issues/PM-1"}})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    try:
        assert [cell.value for cell in workbook["WBS"][1]][12:] == ["Display Number", "Dependency Summary"]
        assert [cell.value for cell in workbook["Dependencies"][1]][5:] == ["Task Title", "Predecessor Title"]
        assert workbook["WBS"]["M2"].value == "1"
        assert workbook["WBS"]["M3"].value == "1.1"
        assert workbook["WBS"]["M4"].value == "1.2"
        assert workbook["WBS"]["N3"].value == "T-ZZZZZZZZ: Foundation (requires_code)"
        assert workbook["Dependencies"]["F2"].value == "=SUM(1,1)"
        assert workbook["Dependencies"]["F2"].data_type == "s"
        assert workbook["Dependencies"]["G2"].value == "Foundation"
        assert workbook["Traceability"]["D2"].value == "/api/v1/issues/PM-1"
        for sheet, cells in {"WBS": ("L3", "M3", "N3"), "Dependencies": ("F2", "G2"), "Traceability": ("D2",)}.items():
            for coordinate in cells:
                assert workbook[sheet][coordinate].fill.fgColor == workbook["WBS"]["B3"].fill.fgColor
                workbook[sheet][coordinate] = "Derived edit is ignored"
        output = io.BytesIO()
        workbook.save(output)
    finally:
        workbook.close()
    candidate = exchange.preview_import(output.getvalue(), change, change)
    assert candidate.changes == candidate.conflicts == candidate.findings == ()
    workbook = openpyxl.load_workbook(io.BytesIO(output.getvalue()))
    workbook["WBS"]["M3"] = "=1+1"
    unsafe = io.BytesIO()
    workbook.save(unsafe)
    workbook.close()
    with pytest.raises(PmError, match="formula cell"):
        exchange.preview_import(unsafe.getvalue(), change, change)


def test_derived_numbering_uses_parent_order_and_bounds_summary_text() -> None:
    exchange = WorkbookExchange(max_text_chars=128)
    change = _change()
    document = change.documents[0]
    source = document.source.replace("order: 10\ntype: Backend", "order: 30\ntype: Backend", 1)
    title = ("@" + "Long foundation " * 7).rstrip()
    source = source.replace("- [ ] Foundation", "- [ ] " + title)
    change = replace(change, documents=(replace(document, source=source),))
    payload = exchange.export(change, {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    try:
        assert workbook["WBS"]["M3"].value == "1.2"
        assert workbook["WBS"]["M4"].value == "1.1"
        assert len(workbook["WBS"]["N3"].value) <= 128
        assert workbook["WBS"]["N3"].value.endswith("See Dependencies.")
        assert workbook["Dependencies"]["G2"].value == title
        assert workbook["Dependencies"]["G2"].data_type == "s"
    finally:
        workbook.close()
    assert exchange.preview_import(payload.data, change, change).changes == ()


def test_deep_derived_numbering_does_not_recurse_and_remains_importable() -> None:
    change = _change()
    document = change.documents[0]
    group = document.source.split("<!-- symphony-task:", 1)[0]
    group_block = group[group.index("<!-- symphony-group:"):]
    foundation = document.source.split("<!-- symphony-task:T-ZZZZZZZZ -->", 1)[1]
    blocks = []
    parent = "G-ABCDEFGH"
    for index in range(1_100):
        group_id = f"G-{index:08d}"
        blocks.append(group_block.replace("G-ABCDEFGH", group_id).replace("parent_id: null", "parent_id: " + parent))
        parent = group_id
    leaf = "<!-- symphony-task:T-ZZZZZZZZ -->" + foundation.replace("parent_id: G-ABCDEFGH", "parent_id: " + parent)
    change = replace(change, documents=(replace(document, source=group + "".join(blocks) + leaf),))
    exchange = WorkbookExchange(max_text_chars=128)
    payload = exchange.export(change, {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data), read_only=True)
    try:
        number = workbook["WBS"]["M1103"].value
        assert len(number) <= 128 and number.endswith("See Parent ID.")
    finally:
        workbook.close()
    assert exchange.preview_import(payload.data, change, change).changes == ()


@pytest.mark.parametrize("sheet,column", [("WBS", 15), ("Dependencies", 8)])
def test_derived_columns_keep_exact_worksheet_bounds(sheet, column) -> None:
    exchange = WorkbookExchange()
    change = _change()
    workbook = openpyxl.load_workbook(io.BytesIO(exchange.export(change, {}).data))
    workbook[sheet].cell(1, column, "Unsupported extra column")
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(PmError, match="worksheet dimensions exceed limit"):
        exchange.preview_import(output.getvalue(), change, change)


def test_nondefault_instance_limits_are_enforced() -> None:
    payload = WorkbookExchange().export(_change(), {})
    with pytest.raises(PmError, match="compressed workbook limit"):
        WorkbookExchange(max_workbook_bytes=len(payload.data) - 1).inspect_metadata(payload.data)
    with pytest.raises(PmError, match="uncompressed-size limit"):
        WorkbookExchange(max_uncompressed_bytes=1).inspect_metadata(payload.data)
    with pytest.raises(PmError, match="ZIP entry limit"):
        WorkbookExchange(max_zip_entries=1).inspect_metadata(payload.data)
    with pytest.raises(PmError) as exc:
        WorkbookExchange(max_nodes=2).export(_change(), {})
    assert exc.value.details["findings"] == ["node_limit_exceeded"]
    with pytest.raises(PmError) as exc:
        WorkbookExchange(max_edges=0).export(_change(), {})
    assert "edge_limit_exceeded" in exc.value.details["findings"]
    with pytest.raises(PmError) as exc:
        WorkbookExchange(max_text_chars=3).export(_change(), {})
    assert "text_limit_exceeded" in exc.value.details["findings"]


def test_print_html_is_readable_and_escapes_active_markup() -> None:
    change = _change("Readable")
    document = change.documents[0]
    source = "# Heading\n\nA paragraph.\n\n- One\n- Two\n\n<script>alert(1)</script>\n"
    change = ChangeView(**{**change.__dict__, "documents": (DocumentView(document.document_id, document.path, source, document.content_hash, document.state, document.editable),)})
    exported = _render_export(change, "html")
    text = exported.data.decode()
    assert "<h1>Heading</h1>" in text
    assert "<ul><li>One</li><li>Two</li></ul>" in text
    assert "<script>" not in text and "&lt;script&gt;alert(1)&lt;/script&gt;" in text


def test_three_way_conflict_is_explicit() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change("base"), {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    workbook["WBS"]["F3"] = "imported"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    candidate = exchange.preview_import(
        output.getvalue(), _change("base"), _change("current", "sha256:current")
    )
    assert [(item.item_id, item.field) for item in candidate.conflicts] == [
        ("T-ABCDEFGH", "title")
    ]


def test_formula_external_relationship_and_zip_bomb_are_rejected() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change(), {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    workbook["WBS"]["A2"] = "=1+1"
    formula = io.BytesIO()
    workbook.save(formula)
    workbook.close()
    with pytest.raises(PmError, match="formula cell"):
        exchange.preview_import(formula.getvalue(), _change(), _change())

    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"x" * (MAX_UNCOMPRESSED + 1))
    with pytest.raises(PmError, match="uncompressed-size"):
        exchange.preview_import(bomb.getvalue(), _change(), _change())


@pytest.mark.parametrize(
    ("member", "replacement"),
    [
        ("_rels/.rels", b"<Relationships><broken>"),
        ("xl/worksheets/sheet2.xml", b"<worksheet><broken>"),
        ("_rels/.rels", b'<!DOCTYPE x [<!ENTITY y "z">]><Relationships/>'),
    ],
)
def test_malformed_or_entity_xml_is_a_stable_unsafe_workbook(
    member: str, replacement: bytes,
) -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change(), {})
    hostile = _rewrite_member(payload.data, member, lambda _value: replacement)
    with pytest.raises(PmError) as exc_info:
        exchange.preview_import(hostile, _change(), _change())
    assert exc_info.value.code == "unsafe_workbook"


def test_duplicate_task_hidden_beyond_declared_dimension_is_rejected() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change(), {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    sheet = workbook["WBS"]
    sheet.append([cell.value for cell in sheet[2]])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    def hide_last_row(xml: bytes) -> bytes:
        return re.sub(br'<dimension ref="[^"]+"', b'<dimension ref="A1:L3"', xml)

    hostile = _rewrite_member(
        output.getvalue(), "xl/worksheets/sheet2.xml", hide_last_row
    )
    with pytest.raises(PmError, match="duplicate task ID"):
        exchange.preview_import(hostile, _change(), _change())


def test_csv_neutralizes_formula_prefix_and_preserves_unicode() -> None:
    data = WorkbookExchange().render_csv(_change("@위험"), excel_utf8=True)
    text = data.decode("utf-8-sig")
    assert "'@위험" in text


def test_import_covers_task_details_and_new_task_fields_with_temp_references() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change(), {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    workbook["Task Details"]["B3"] = "Imported description"
    workbook["Task Details"]["C3"] = "One\nTwo"
    workbook["WBS"].append(("Add", "NEW-1", "G-ABCDEFGH", 20, "Backend", "New task", "Owner", "codex", "high", "2026-10-10", "2d", ""))
    workbook["Task Details"].append(("NEW-1", "New description", "Accepted", "Run tests", "API", "Low", "Note"))
    workbook["Dependencies"].append(("NEW-1", "T-ABCDEFGH", "requires_code", ""))
    workbook["Traceability"].append(("NEW-1", "requirement", "REQ-2", ""))
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    candidate = exchange.preview_import(output.getvalue(), _change(), _change())
    assert {(change.item_id, change.field) for change in candidate.changes} >= {
        ("T-ABCDEFGH", "description"),
        ("T-ABCDEFGH", "acceptance"),
    }
    assigned = candidate.assigned_ids["NEW-1"]
    addition = next(change for change in candidate.changes if change.action == "add")
    assert addition.item_id == assigned
    assert addition.imported["description"] == "New description"
    assert addition.imported["dependencies"][0].task_id == "T-ABCDEFGH"
    assert addition.imported["requirement_refs"] == ("REQ-2",)


def test_optional_sheet_omission_retains_content_and_version_is_strict() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change(), {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    workbook.remove(workbook["Dependencies"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    candidate = exchange.preview_import(output.getvalue(), _change(), _change())
    assert not any(change.field == "dependencies" for change in candidate.changes)

    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    workbook["Metadata"]["A2"] = "999"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(PmError, match="format_version"):
        exchange.preview_import(output.getvalue(), _change(), _change())


def test_missing_dependency_and_trace_rows_retain_until_explicit_delete() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change(), {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    workbook["Dependencies"].delete_rows(2)
    workbook["Traceability"].delete_rows(2)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    candidate = exchange.preview_import(output.getvalue(), _change(), _change())
    assert not any(change.field in {"dependencies", "requirement_refs"} for change in candidate.changes)

    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    workbook["Dependencies"]["E2"] = "Delete"
    workbook["Traceability"]["E2"] = "Delete"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    candidate = exchange.preview_import(output.getvalue(), _change(), _change())
    fields = {change.field: change.imported for change in candidate.changes if change.item_id == "T-ABCDEFGH"}
    assert fields["dependencies"] == ()
    assert fields["requirement_refs"] == ()


def test_delete_conflicts_when_affected_task_changed_since_export() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change("base"), {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload.data))
    workbook["WBS"]["A3"] = "Delete"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    candidate = exchange.preview_import(
        output.getvalue(), _change("base"), _change("current", "sha256:current")
    )
    assert [(item.item_id, item.field) for item in candidate.conflicts] == [
        ("T-ABCDEFGH", "__delete__")
    ]


def test_declared_full_grid_dimension_is_rejected_before_iteration() -> None:
    exchange = WorkbookExchange()
    payload = exchange.export(_change(), {})
    source = io.BytesIO(payload.data)
    output = io.BytesIO()
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(output, "w") as outgoing:
        for info in incoming.infolist():
            data = incoming.read(info)
            if info.filename == "xl/worksheets/sheet2.xml":
                data = re.sub(b'ref="A1:[A-Z]+[0-9]+"', b'ref="A1:XFD1048576"', data, count=1)
            outgoing.writestr(info, data)
    with pytest.raises(PmError, match="dimensions exceed"):
        exchange.preview_import(output.getvalue(), _change(), _change())
